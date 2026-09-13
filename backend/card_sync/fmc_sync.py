"""
FMC PO Summary Cards Synchronization Module
Sri Balaji Traders Automation System
"""

from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from .constants import (
    REGULAR_FONT,
    BOLD_FONT,
    RED_BOLD_FONT,
    GREEN_BOLD_FONT,
    THIN_BORDER,
    normalize_activity,
    match_activity_to_column,
)


def is_fmc_row_occupied(ws, d_r, amount_col):
    """
    Checks if a row in FMC card already contains valid data.
    """
    c4 = str(ws.cell(d_r, 4).value or '').strip()
    c8 = str(ws.cell(d_r, 8).value or '').strip()
    c10 = str(ws.cell(d_r, 10).value or '').strip()
    if c4 or c8 or c10:
        return True

    for c in range(12, amount_col):
        v = ws.cell(d_r, c).value
        if isinstance(v, (int, float)) and v > 0:
            return True

    return False


def write_sheet_end_summary_table(ws, start_row=240):
    """
    Writes the Sheet Summary Table at row 240 (Cols H to N) summarizing
    all POs, products, crops, activities, budgets (=P{r}), spent (=S{r}), and balances (=T{r}) on this sheet.
    """
    header_font = BOLD_FONT
    regular_font = REGULAR_FONT
    bold_font = BOLD_FONT

    entries = []
    for b in range(11):
        r = 1 + b * 19
        if r > ws.max_row:
            break
        po_val = ws.cell(r, 1).value
        if not po_val or not str(po_val).strip().startswith('500'):
            continue
        po_str = str(po_val).strip()
        prod = str(ws.cell(r, 7).value or '').strip()
        crop1 = str(ws.cell(r, 9).value or '').strip()
        crop2 = str(ws.cell(r + 1, 9).value or '').strip()

        for idx in range(8):
            curr_r = r + 8 + idx
            act_val = ws.cell(curr_r, 15).value
            if act_val and str(act_val).strip() not in ['0', 'None', '']:
                act_str = str(act_val).strip()
                c_crop = crop1
                if idx == 1 and crop2:
                    c_crop = crop2
                entries.append({
                    'po': po_str,
                    'prod': prod,
                    'crop': c_crop,
                    'act': act_str,
                    'budget_ref': f'P{curr_r}',
                    'spent_ref': f'S{curr_r}',
                    'bal_ref': f'T{curr_r}'
                })

    if not entries:
        return

    # Clear prior table area (rows start_row to start_row + 45, cols 8 to 14)
    for cr in range(start_row, start_row + 45):
        for cc in range(8, 15):
            ws.cell(cr, cc).value = None
            ws.cell(cr, cc).border = Border()

    # Headers in row start_row
    headers = ['Pos', 'Product', 'Crop', 'Activity', 'Budget', 'Spent', 'Balance']
    for c_i, h_text in enumerate(headers, start=8):
        c = ws.cell(start_row, c_i, h_text)
        c.font = header_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='right' if h_text in ['Budget', 'Spent', 'Balance'] else 'center', vertical='center')

    # Data rows
    for e_idx, e in enumerate(entries):
        row_r = start_row + 1 + e_idx

        c = ws.cell(row_r, 8, e['po'])
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='center', vertical='center')

        c = ws.cell(row_r, 9, e['prod'])
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='center', vertical='center')

        c = ws.cell(row_r, 10, e['crop'])
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='center', vertical='center')

        c = ws.cell(row_r, 11, e['act'])
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='center', vertical='center')

        c = ws.cell(row_r, 12, f"={e['budget_ref']}")
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='right', vertical='center')
        c.number_format = '#,##0'

        c = ws.cell(row_r, 13, f"={e['spent_ref']}")
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='right', vertical='center')
        c.number_format = '#,##0.00'

        c = ws.cell(row_r, 14, f"={e['bal_ref']}")
        c.font = regular_font
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal='right', vertical='center')
        c.number_format = '#,##0.00'

    # TOTAL row
    tot_r = start_row + 1 + len(entries)
    c = ws.cell(tot_r, 8, 'TOTAL')
    c.font = bold_font
    c.border = THIN_BORDER
    c.alignment = Alignment(horizontal='center', vertical='center')

    for cc in [9, 10, 11]:
        c = ws.cell(tot_r, cc, '')
        c.border = THIN_BORDER

    c = ws.cell(tot_r, 12, f"=SUM(L{start_row+1}:L{tot_r-1})")
    c.font = bold_font
    c.border = THIN_BORDER
    c.alignment = Alignment(horizontal='right', vertical='center')
    c.number_format = '#,##0'

    c = ws.cell(tot_r, 13, f"=SUM(M{start_row+1}:M{tot_r-1})")
    c.font = bold_font
    c.border = THIN_BORDER
    c.alignment = Alignment(horizontal='right', vertical='center')
    c.number_format = '#,##0.00'

    c = ws.cell(tot_r, 14, f"=SUM(N{start_row+1}:N{tot_r-1})")
    c.font = bold_font
    c.border = THIN_BORDER
    c.alignment = Alignment(horizontal='right', vertical='center')
    c.number_format = '#,##0.00'


def sync_fmc_cards_workbook(wb_cards, tbm_data, service_charge_percent=5.0):
    """
    Synchronizes FMC PO Summary Cards (where cards are stacked in blocks of 19 rows across sheets).
    Appends new spending entries into available rows below existing ones (non-destructive),
    updates bottom totals, right summary tables, and Sheet1 master overview table.
    """
    sv_rate = float(service_charge_percent) / 100.0
    card_links = {}  # maps (po_number, normalized_activity) -> (sheet_name, summary_row_idx)
    updated_cards_count = 0

    for sname in wb_cards.sheetnames:
        if sname.lower() == 'sheet1':
            continue
        ws = wb_cards[sname]

        for b in range(11):
            r = 1 + b * 19
            if r > ws.max_row:
                break

            po_val = str(ws.cell(r, 1).value or '').strip().upper()
            if not po_val or not po_val.startswith('500'):
                continue

            prod_card = str(ws.cell(r, 7).value or '').strip()
            crop1_card = str(ws.cell(r, 9).value or '').strip()
            crop2_card = str(ws.cell(r + 1, 9).value or '').strip()
            crop_card = f"{crop1_card} / {crop2_card}".strip(" /") if crop2_card else crop1_card
            area_card = str(ws.cell(r + 2, 7).value or '').strip()
            am_card = str(ws.cell(r + 3, 4).value or '').strip()
            prefix = po_val[3:5] if len(po_val) >= 5 else 'BB'
            budget_type = 'Brand' if prefix == 'BB' else 'Promo'

            # 1. Identify Activity Columns in row (r + 5)
            act_col_map = {}
            last_act_col = 12
            for c in range(12, 16):
                h_val = ws.cell(r + 5, c).value
                if h_val and str(h_val).strip().lower() == 'amount':
                    last_act_col = c
                    break
                elif h_val:
                    act_col_map[normalize_activity(h_val)] = c
                    last_act_col = c + 1

            amount_col = last_act_col

            # 2. Non-Destructive Append in Main Table Data Rows (rows r + 7 to r + 15)
            first_data_r = r + 7
            last_data_r = r + 15
            table_tot_r = r + 17

            curr_data_r = first_data_r
            while curr_data_r <= last_data_r and is_fmc_row_occupied(ws, curr_data_r, amount_col):
                curr_data_r += 1

            tbm_rows_for_po = tbm_data.get(po_val, [])
            newly_added = 0

            for t_item in tbm_rows_for_po:
                if curr_data_r > last_data_r:
                    break

                ws.cell(curr_data_r, 1).value = None
                ws.cell(curr_data_r, 2).value = None
                ws.cell(curr_data_r, 3, area_card).font = REGULAR_FONT
                ws.cell(curr_data_r, 4, po_val).font = REGULAR_FONT
                ws.cell(curr_data_r, 5, budget_type).font = REGULAR_FONT
                ws.cell(curr_data_r, 6, t_item.get('product') or prod_card).font = REGULAR_FONT
                ws.cell(curr_data_r, 7, t_item.get('crop') or crop_card).font = REGULAR_FONT
                ws.cell(curr_data_r, 8, t_item.get('activity', '')).font = REGULAR_FONT
                ws.cell(curr_data_r, 9, am_card).font = REGULAR_FONT
                ws.cell(curr_data_r, 10, t_item.get('tbm', '')).font = REGULAR_FONT
                ws.cell(curr_data_r, 11, t_item.get('num_activities', 1)).font = REGULAR_FONT

                target_col = match_activity_to_column(t_item.get('activity', ''), act_col_map)
                if not target_col or target_col >= amount_col:
                    target_col = 12

                for c in range(12, amount_col):
                    if c == target_col:
                        amt_cell = ws.cell(curr_data_r, c, t_item.get('total_amount', 0.0))
                        amt_cell.font = REGULAR_FONT
                        amt_cell.number_format = '#,##0'
                    else:
                        ws.cell(curr_data_r, c, None)

                act_cell_refs = [f"{get_column_letter(c)}{curr_data_r}" for c in range(12, amount_col)]
                amt_formula = f"=SUM({','.join(act_cell_refs)})" if act_cell_refs else 0
                c_amt = ws.cell(curr_data_r, amount_col, amt_formula)
                c_amt.font = REGULAR_FONT
                c_amt.number_format = '#,##0'

                curr_data_r += 1
                newly_added += 1

            if newly_added > 0:
                updated_cards_count += 1

            # 3. Bottom Totals row for main table (row r + 17)
            for c in range(12, amount_col):
                cl = get_column_letter(c)
                c_tot = ws.cell(table_tot_r, c, f"=SUM({cl}{first_data_r}:{cl}{last_data_r})")
                c_tot.font = RED_BOLD_FONT
                c_tot.number_format = '#,##0'

            amt_cl = get_column_letter(amount_col)
            c_amt_tot = ws.cell(table_tot_r, amount_col, f"=SUM({amt_cl}{first_data_r}:{amt_cl}{last_data_r})")
            c_amt_tot.font = RED_BOLD_FONT
            c_amt_tot.number_format = '#,##0'

            # 4. Right Summary Table (Rows r + 8 to r + 15)
            first_sum_r = r + 8
            last_sum_r = r + 15
            sum_tot_r = r + 16

            for idx in range(8):
                curr_r = first_sum_r + idx
                act_name_val = ws.cell(curr_r, 15).value

                if act_name_val and str(act_name_val).strip() not in ['0', 'None', '']:
                    norm_act = normalize_activity(act_name_val)
                    target_col = match_activity_to_column(act_name_val, act_col_map)
                    if not target_col:
                        target_col = 12
                    t_cl = get_column_letter(target_col)

                    # SPENT -> points to bottom table total for this activity
                    ws.cell(curr_r, 17, f"={t_cl}{table_tot_r}").font = REGULAR_FONT
                    ws.cell(curr_r, 17).number_format = '#,##0'

                    # SV Charges -> =SPENT * sv_rate
                    ws.cell(curr_r, 18, f"=Q{curr_r}*{sv_rate}").font = REGULAR_FONT
                    ws.cell(curr_r, 18).number_format = '#,##0.00'

                    # TOTAL IV -> =SPENT + SV Charges
                    ws.cell(curr_r, 19, f"=Q{curr_r}+R{curr_r}").font = REGULAR_FONT
                    ws.cell(curr_r, 19).number_format = '#,##0.00'

                    # BALANCE -> =BUDGET - TOTAL IV
                    ws.cell(curr_r, 20, f"=P{curr_r}-S{curr_r}").font = REGULAR_FONT
                    ws.cell(curr_r, 20).number_format = '#,##0.00'

                    card_links[(po_val, norm_act)] = (sname, curr_r)
                else:
                    ws.cell(curr_r, 17, 0).font = REGULAR_FONT
                    ws.cell(curr_r, 18, 0).font = REGULAR_FONT
                    ws.cell(curr_r, 19, 0).font = REGULAR_FONT
                    ws.cell(curr_r, 20, f"=P{curr_r}-S{curr_r}").font = REGULAR_FONT

            # Total row in Right Summary Table
            ws.cell(sum_tot_r, 15, "TOTAL").font = RED_BOLD_FONT

            c_sum_b = ws.cell(sum_tot_r, 16, f"=SUM(P{first_sum_r}:P{last_sum_r})")
            c_sum_b.font = GREEN_BOLD_FONT
            c_sum_b.number_format = '#,##0'

            c_sum_sp = ws.cell(sum_tot_r, 17, f"=SUM(Q{first_sum_r}:Q{last_sum_r})")
            c_sum_sp.font = REGULAR_FONT
            c_sum_sp.number_format = '#,##0'

            c_sum_sv = ws.cell(sum_tot_r, 18, f"=SUM(R{first_sum_r}:R{last_sum_r})")
            c_sum_sv.font = REGULAR_FONT
            c_sum_sv.number_format = '#,##0.00'

            c_sum_iv = ws.cell(sum_tot_r, 19, f"=SUM(S{first_sum_r}:S{last_sum_r})")
            c_sum_iv.font = REGULAR_FONT
            c_sum_iv.number_format = '#,##0.00'

            c_sum_bal = ws.cell(sum_tot_r, 20, f"=P{sum_tot_r}-S{sum_tot_r}")
            c_sum_bal.font = RED_BOLD_FONT
            c_sum_bal.number_format = '#,##0.00'

        # 5. Write / Update Sheet End Summary Table at Row 240 (Cols H to N)
        write_sheet_end_summary_table(ws, start_row=240)

    # 6. Update Sheet1 Master Overview Table
    if 'Sheet1' in wb_cards.sheetnames:
        ws1 = wb_cards['Sheet1']
        curr_po = ''

        for r in range(4, ws1.max_row + 1):
            val_c6 = ws1.cell(r, 6).value
            if val_c6 and str(val_c6).strip().lower() == 'total':
                ws1.cell(r, 7, f"=SUM(G4:G{r-1})")
                ws1.cell(r, 8, f"=SUM(H4:H{r-1})")
                ws1.cell(r, 9, f"=G{r}-H{r}")
                break

            po_cell = ws1.cell(r, 3).value
            if po_cell and str(po_cell).strip():
                curr_po = str(po_cell).strip().upper()

            act_cell = ws1.cell(r, 6).value
            if act_cell and str(act_cell).strip():
                norm_act = normalize_activity(act_cell)
                key = (curr_po, norm_act)

                if key in card_links:
                    card_sheet, card_row = card_links[key]
                    ws1.cell(r, 8, f"='{card_sheet}'!S{card_row}").number_format = '#,##0.00'
                else:
                    ws1.cell(r, 8, 0).number_format = '#,##0.00'

                ws1.cell(r, 9, f"=G{r}-H{r}").number_format = '#,##0.00'

    return updated_cards_count
