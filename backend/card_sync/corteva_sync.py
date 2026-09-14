"""
Corteva PO Summary Cards Synchronization Module
Sri Balaji Traders Automation System
"""

from openpyxl.utils import get_column_letter
from .constants import (
    REGULAR_FONT,
    BOLD_FONT,
    RED_BOLD_FONT,
    normalize_activity,
    match_activity_to_column,
)

try:
    from common.formatters import normalize_po
except ImportError:
    import sys
    from pathlib import Path
    sys.path.append(str(Path(__file__).resolve().parent.parent))
    from common.formatters import normalize_po


def resolve_cell_text(ws, cell_val):
    """
    Dereferences formula references like '=T3', '=$T$3' to get the actual text value.
    Returns empty string if referenced cell is empty or invalid.
    """
    if cell_val is None:
        return ""
    s = str(cell_val).strip()
    if s.startswith('='):
        clean_ref = s.lstrip('=').strip().replace('$', '')
        try:
            val = ws[clean_ref].value
            if val is not None:
                s_val = str(val).strip()
                if s_val.startswith('='):
                    clean_ref2 = s_val.lstrip('=').strip().replace('$', '')
                    val = ws[clean_ref2].value
                if val is not None and str(val).strip() not in ['0', 'None', 'none', '']:
                    return str(val).strip()
            return ""  # Empty or invalid reference
        except Exception:
            return ""
    return s


def is_corteva_row_occupied(ws, r, amount_col):
    """
    Checks if a row in Corteva card already contains valid data.
    """
    # Check key text columns: Col 4 (PO Number), Col 8 (Activity), Col 10 (TBM)
    c4 = str(ws.cell(r, 4).value or '').strip()
    c8 = str(ws.cell(r, 8).value or '').strip()
    c10 = str(ws.cell(r, 10).value or '').strip()
    if c4 or c8 or c10:
        return True

    # Check if any activity column (12 to amount_col - 1) has a numeric value > 0
    for c in range(12, amount_col):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)) and v > 0:
            return True

    return False


def sync_corteva_cards_workbook(wb_cards, tbm_data, service_charge_percent=5.0):
    """
    Synchronizes Corteva PO Summary Cards (where each product is on its own worksheet).
    Appends new spending entries into available rows below existing ones (non-destructive),
    and updates bottom sums, Right Summary table, and overview sheets.
    Always maintains bottom sums and Right Summary table formulas on ALL product card sheets,
    even if no new spending rows were added in this run.
    """
    sv_rate = float(service_charge_percent) / 100.0
    updated_cards_count = 0
    card_activity_links = {}  # (po_val, norm_act) -> (sheet_name, summary_row_idx)

    for name in wb_cards.sheetnames:
        if name.strip().lower() in ['sheet1', 'processed_emails', 'pos & expenditures', 'overview', 'summary']:
            continue

        ws = wb_cards[name]
        po_val = str(ws.cell(6, 1).value or '').strip().upper()
        if not po_val:
            po_val = str(ws.cell(10, 4).value or '').strip().upper()
        if not po_val:
            for r_chk in range(12, 16):
                val_c4 = str(ws.cell(r_chk, 4).value or '').strip().upper()
                if val_c4 and val_c4.startswith('48'):
                    po_val = val_c4
                    break

        norm_po_val = normalize_po(po_val)

        # 1. Locate Amount column in Row 10 / 11
        amount_col = 19
        for c in range(12, 23):
            val_10 = str(ws.cell(10, c).value or '').strip().lower()
            val_11 = str(ws.cell(11, c).value or '').strip().lower()
            if val_10 == 'amount' or val_11 == 'amount':
                amount_col = c
                break

        # Map activity names in row 10 / 11 headers: norm_act -> col_idx
        # Priority 1: Headers in Row 10 / 11
        act_norm_to_col = {}
        for c in range(12, amount_col):
            h_val = ws.cell(10, c).value or ws.cell(11, c).value
            resolved_h = resolve_cell_text(ws, h_val)
            if resolved_h and str(resolved_h).strip() not in ['0', 'None', '', 'none'] and resolved_h.lower() != 'amount':
                norm = normalize_activity(resolved_h)
                if norm and norm.lower() != 'amount':
                    act_norm_to_col[norm] = c

        # Priority 2: Rate grid at T3:T9 (which corresponds 1:1 with columns 12..)
        for idx, r_grid in enumerate(range(3, 10)):
            c = 12 + idx
            if c < amount_col:
                grid_act = resolve_cell_text(ws, ws.cell(r_grid, 20).value)
                if grid_act and str(grid_act).strip().upper() not in ['TOTAL', '0', 'NONE', '']:
                    norm_grid = normalize_activity(grid_act)
                    if norm_grid and norm_grid not in act_norm_to_col and norm_grid.lower() != 'amount':
                        act_norm_to_col[norm_grid] = c

        # Priority 3: Right Summary table rows 16 to 21
        for r_sum in range(16, 22):
            c = 12 + (r_sum - 16)
            if c < amount_col:
                sum_act = resolve_cell_text(ws, ws.cell(r_sum, 20).value)
                if sum_act and str(sum_act).strip().upper() not in ['TOTAL', '0', 'NONE', '']:
                    norm_sum = normalize_activity(sum_act)
                    if norm_sum and norm_sum not in act_norm_to_col and norm_sum.lower() != 'amount':
                        act_norm_to_col[norm_sum] = c

        # 2. Append new spending rows from tbm_data if available
        tbm_rows_for_po = tbm_data.get(po_val, [])
        if not tbm_rows_for_po and norm_po_val:
            for k, rows in tbm_data.items():
                if normalize_po(k) == norm_po_val:
                    tbm_rows_for_po = rows
                    break

        newly_added_count = 0
        if tbm_rows_for_po:
            curr_row = 12
            while curr_row <= 28 and is_corteva_row_occupied(ws, curr_row, amount_col):
                curr_row += 1

            area_val = str(ws.cell(7, 4).value or '').split('-')[-1].strip()
            if not area_val:
                area_val = str(ws.cell(9, 2).value or '').split('-')[-1].strip()

            for t_item in tbm_rows_for_po:
                if curr_row > 28:
                    break  # Card reached maximum row capacity (row 28)

                ws.cell(curr_row, 1).value = None  # I.V NO
                ws.cell(curr_row, 2).value = None  # DATE
                ws.cell(curr_row, 3, area_val or t_item.get('area', '')).font = REGULAR_FONT
                ws.cell(curr_row, 4, po_val).font = REGULAR_FONT
                ws.cell(curr_row, 5, "Marketing").font = REGULAR_FONT
                ws.cell(curr_row, 6, t_item.get('product') or name).font = REGULAR_FONT
                ws.cell(curr_row, 7, t_item.get('crop') or "All Crops").font = REGULAR_FONT
                ws.cell(curr_row, 8, t_item.get('activity', '')).font = REGULAR_FONT
                ws.cell(curr_row, 9, t_item.get('zdgm') or "ZDGM").font = REGULAR_FONT
                ws.cell(curr_row, 10, t_item.get('tbm', '')).font = REGULAR_FONT
                ws.cell(curr_row, 11, t_item.get('num_activities', 1)).font = REGULAR_FONT

                target_col = match_activity_to_column(t_item.get('activity', ''), act_norm_to_col)
                if not target_col or target_col >= amount_col:
                    target_col = 12

                for c in range(12, amount_col):
                    if c == target_col:
                        ws.cell(curr_row, c, t_item.get('total_amount', 0.0)).font = REGULAR_FONT
                        ws.cell(curr_row, c).number_format = '#,##0'
                    else:
                        ws.cell(curr_row, c, None)

                # Row total formula
                parts = [f"{get_column_letter(c)}{curr_row}" for c in range(12, amount_col)]
                ws.cell(curr_row, amount_col, f"=SUM({','.join(parts)})").font = REGULAR_FONT
                ws.cell(curr_row, amount_col).number_format = '#,##0'

                curr_row += 1
                newly_added_count += 1

            if newly_added_count > 0:
                updated_cards_count += 1

        # 3. Ensure all populated data rows (12 to 28) have row sum formula in amount_col
        for r_check in range(12, 29):
            if any(ws.cell(r_check, c).value for c in range(12, amount_col)):
                parts_check = [f"{get_column_letter(c)}{r_check}" for c in range(12, amount_col)]
                ws.cell(r_check, amount_col, f"=SUM({','.join(parts_check)})").font = REGULAR_FONT
                ws.cell(r_check, amount_col).number_format = '#,##0'

        # 4. Bottom Totals in row 29 (sums rows 12 to 28)
        for c in range(12, amount_col):
            cl = get_column_letter(c)
            cell = ws.cell(29, c, f"=SUM({cl}12:{cl}28)")
            cell.font = RED_BOLD_FONT
            cell.number_format = '#,##0'

        amt_cl = get_column_letter(amount_col)
        cell = ws.cell(29, amount_col, f"=SUM({amt_cl}12:{amt_cl}28)")
        cell.font = RED_BOLD_FONT
        cell.number_format = '#,##0'

        # 5. Right Summary block T16:Y (Columns T to Y / 20 to 25)
        # Find the real TOTAL row (search downwards from 25 down to 17)
        tot_row = None
        for r in range(25, 16, -1):
            if str(ws.cell(r, 20).value or '').strip().upper() == 'TOTAL':
                tot_row = r
                break
        if tot_row is None:
            tot_row = 16 + (amount_col - 12)
        if tot_row < 17:
            tot_row = 17

        for r in range(16, tot_row):
            c = 12 + (r - 16)
            if c >= amount_col:
                break
            cl = get_column_letter(c)
            val_t = ws.cell(r, 20).value

            # If intermediate row had a bogus 'TOTAL' from previous buggy runs, clear it
            if str(val_t or '').strip().upper() == 'TOTAL':
                hdr_c = ws.cell(10, c).value or ws.cell(11, c).value
                resolved_hdr = resolve_cell_text(ws, hdr_c)
                if resolved_hdr and str(resolved_hdr).strip() not in ['0', 'None', 'none', '']:
                    val_t = resolved_hdr
                    ws.cell(r, 20, val_t).font = REGULAR_FONT
                else:
                    val_t = 0
                    ws.cell(r, 20, 0).font = REGULAR_FONT
                    ws.cell(r, 21, 0).font = REGULAR_FONT

            # Determine activity name
            act_name = resolve_cell_text(ws, val_t)
            if not act_name:
                hdr_c = ws.cell(10, c).value or ws.cell(11, c).value
                act_name = resolve_cell_text(ws, hdr_c)
                if act_name and val_t is None:
                    ws.cell(r, 20, act_name).font = REGULAR_FONT

            norm_a = normalize_activity(act_name) if act_name else ''
            if norm_a and norm_a.lower() not in ['0', 'none', 'total']:
                card_activity_links[(po_val, norm_a)] = (name, r)

            # SPENT in Col V (22) points to its respective column's bottom sum at row 29
            ws.cell(r, 22, f"={cl}29").font = REGULAR_FONT
            ws.cell(r, 22).number_format = '#,##0'

            # SV Charges in Col W (23)
            curr_w = ws.cell(r, 23).value
            if norm_a and 'bvc' in norm_a.lower() and (curr_w is None or curr_w == 0 or curr_w == ''):
                ws.cell(r, 23).value = None
            else:
                ws.cell(r, 23, f"=V{r}*{sv_rate:.4f}".rstrip('0').rstrip('.')).font = REGULAR_FONT
                ws.cell(r, 23).number_format = '#,##0.00'

            # TOTAL IV in Col X (24) = SPENT + SV Charges
            ws.cell(r, 24, f"=V{r}+W{r}").font = REGULAR_FONT
            ws.cell(r, 24).number_format = '#,##0.00'

            # BALANCE in Col Y (25) = BUDGET - TOTAL IV
            ws.cell(r, 25, f"=U{r}-X{r}").font = REGULAR_FONT
            ws.cell(r, 25).number_format = '#,##0.00'

        # Write TOTAL formulas in tot_row
        ws.cell(tot_row, 20, "TOTAL").font = RED_BOLD_FONT
        for col_idx in [21, 22, 23, 24, 25]:
            cl = get_column_letter(col_idx)
            cell = ws.cell(tot_row, col_idx, f"=SUM({cl}16:{cl}{tot_row-1})")
            cell.font = RED_BOLD_FONT
            cell.number_format = '#,##0.00' if col_idx in [23, 24, 25] else '#,##0'

    # 6. Link overview sheets if present
    for sname in wb_cards.sheetnames:
        if sname.strip().lower() in ['pos & expenditures', 'sheet1']:
            ws_ov = wb_cards[sname]
            curr_po = ''
            for r in range(4, ws_ov.max_row + 1):
                po_cell_val = ws_ov.cell(r, 3).value
                if po_cell_val and str(po_cell_val).strip():
                    curr_po = str(po_cell_val).strip().upper()

                act_cell_val = ws_ov.cell(r, 8).value or ws_ov.cell(r, 6).value
                if act_cell_val and str(act_cell_val).strip():
                    norm_a = normalize_activity(act_cell_val)
                    key = (curr_po, norm_a)
                    if key in card_activity_links:
                        c_sheet, c_row = card_activity_links[key]
                        # Col 10 (J): Spent -> Link to card's TOTAL IV (Col X)
                        ws_ov.cell(r, 10, f"='{c_sheet}'!X{c_row}").number_format = '#,##0.00'
                        # Col 11 (K): Balance = Budget (Col 9 / I) - Spent
                        ws_ov.cell(r, 11, f"=I{r}-J{r}").number_format = '#,##0.00'

    return updated_cards_count
