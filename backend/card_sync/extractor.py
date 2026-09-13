"""
Card Sync Extractor & Status Marker Module
Sri Balaji Traders Automation System
"""

import os
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

from .constants import (
    STATUS_HDR_FONT,
    STATUS_VAL_FONT,
    STATUS_BORDER,
    CENTER_ALIGN,
    normalize_activity,
)

# Import reader from backend.tbm
try:
    from tbm.reader import extract_activities_from_sheet
    from excel_parser import load_any_workbook
except ImportError:
    import sys
    sys.path.append(str(Path(__file__).resolve().parent.parent))
    from tbm.reader import extract_activities_from_sheet
    from excel_parser import load_any_workbook


def find_summary_files(summary_path):
    """
    Finds all valid Excel summary files given a directory path or single file path.
    Excludes temporary Excel lock files (starting with '~$').
    """
    p = Path(summary_path).resolve()
    if not p.exists():
        raise FileNotFoundError(f"TBM Summary path does not exist: {summary_path}")

    if p.is_file():
        if p.name.startswith("~$") or not p.suffix.lower() in ['.xlsx', '.xlsm', '.xls']:
            return []
        return [p]

    # It's a directory
    files = [
        f for f in p.iterdir()
        if f.is_file()
        and f.suffix.lower() in ['.xlsx', '.xlsm', '.xls']
        and not f.name.startswith("~$")
    ]
    # Sort files deterministically
    return sorted(files, key=lambda x: x.name.lower())


def extract_all_tbms_spent(summary_path, force=False):
    """
    Extracts spent records across one or more All-TBMs summary workbooks.
    Skips rows/tables that are already marked as 'DONE' in Column R (Col 18) unless force=True.
    
    Returns:
        tbm_data: dict of po_number -> list of spending records:
                  [{po, tbm, product, crop, activity, num_activities, total_amount, area, zdgm}, ...]
        items_to_mark: list of dicts with file, sheet, header_row, data_rows for status tracking
        skipped_files_count: count of files skipped because all entries were already DONE
        total_files_scanned: total number of summary files checked
    """
    files = find_summary_files(summary_path)
    if not files:
        raise ValueError(f"No valid Excel files (.xlsx / .xls) found in: {summary_path}")

    tbm_data = {}
    items_to_mark = []
    skipped_files_count = 0
    total_files_scanned = len(files)

    for ef in files:
        try:
            wb = load_any_workbook(ef)
        except Exception as e:
            print(f"Error loading summary file {ef.name}: {e}")
            continue

        file_has_unmarked = False
        file_had_any_data = False

        # 1. Scan activity sheets
        activity_sheets = [
            s for s in wb.sheetnames
            if s.strip().lower() not in ['processed_emails', 'tbm amount summary']
        ]

        file_acts = []
        for sname in activity_sheets:
            ws = wb[sname]
            acts = extract_activities_from_sheet(
                ws, default_tbm_name="All-TBMs", file_name=ef.name, skip_done=not force
            )
            if acts:
                for a in acts:
                    a['_sheet_name'] = sname
                file_acts.extend(acts)

            # Check if sheet had any data at all (to detect fully DONE files)
            if not file_had_any_data:
                raw_acts = extract_activities_from_sheet(
                    ws, default_tbm_name="All-TBMs", file_name=ef.name, skip_done=False
                )
                if raw_acts:
                    file_had_any_data = True

        if file_acts:
            file_has_unmarked = True
            # Group extracted rows by table: (sheet_name, _hr_idx)
            table_groups = {}
            for a in file_acts:
                sname = a.get('_sheet_name')
                hr = a.get('_hr_idx')
                key = (sname, hr)
                if key not in table_groups:
                    table_groups[key] = []
                table_groups[key].append(a)

            for (sname, hr), tbl_rows in table_groups.items():
                first_r = tbl_rows[0]
                po_val = str(first_r.get('po_number', '')).strip().upper()
                if not po_val or po_val in ['NO PO', 'NO_PO', 'NONE']:
                    continue

                tbm_val = first_r.get('tbm', 'TBM')
                prod_val = first_r.get('product', '')
                crop_val = first_r.get('crop', '')
                act_val = first_r.get('activity', '')
                area_val = first_r.get('territory', '')
                zdgm_val = first_r.get('zdgm', '')

                row_indices = [r.get('_row_idx') for r in tbl_rows if r.get('_row_idx')]
                tot_amt = sum(float(r.get('total', 0.0) or 0.0) for r in tbl_rows)

                spend_record = {
                    'po': po_val,
                    'tbm': tbm_val,
                    'product': prod_val,
                    'crop': crop_val,
                    'activity': act_val,
                    'num_activities': len(tbl_rows),
                    'total_amount': tot_amt,
                    'area': area_val,
                    'zdgm': zdgm_val
                }

                if po_val not in tbm_data:
                    tbm_data[po_val] = []
                tbm_data[po_val].append(spend_record)

                items_to_mark.append({
                    'file_path': ef,
                    'sheet_name': sname,
                    'header_row': hr,
                    'data_rows': row_indices,
                    'po': po_val,
                    'tbm': tbm_val,
                    'activity': act_val,
                    'product': prod_val,
                    'crop': crop_val,
                    'is_summary_sheet': False
                })

        elif not file_had_any_data:
            # 2. Fallback: Read from 'TBM Amount Summary' sheet if no activity sheets had tables
            target_sheet = None
            for sname in wb.sheetnames:
                if 'tbm amount summary' in sname.lower() or 'amount summary' in sname.lower():
                    target_sheet = wb[sname]
                    target_sname = sname
                    break

            if target_sheet is not None:
                header_row = 2
                for r in range(1, min(10, target_sheet.max_row + 1)):
                    c1 = str(target_sheet.cell(r, 1).value or '').strip().lower()
                    c2 = str(target_sheet.cell(r, 2).value or '').strip().lower()
                    if 'po' in c1 or 'po' in c2 or 'tbm' in c2:
                        header_row = r
                        break

                for r in range(header_row + 1, target_sheet.max_row + 1):
                    po_val = str(target_sheet.cell(r, 1).value or '').strip().upper()
                    if not po_val or 'TOTAL' in po_val or 'GRAND' in po_val:
                        continue

                    file_had_any_data = True
                    status_val = str(target_sheet.cell(r, 8).value or '').strip().upper()
                    if not force and status_val == 'DONE':
                        continue

                    file_has_unmarked = True
                    tbm_val = str(target_sheet.cell(r, 2).value or '').strip()
                    prod_val = str(target_sheet.cell(r, 3).value or '').strip()
                    crop_val = str(target_sheet.cell(r, 4).value or '').strip()
                    act_val = str(target_sheet.cell(r, 5).value or '').strip()

                    try:
                        num_act_val = int(target_sheet.cell(r, 6).value or 0)
                    except Exception:
                        num_act_val = 1

                    try:
                        amt_val = float(str(target_sheet.cell(r, 7).value or '0').replace(',', '').strip())
                    except Exception:
                        amt_val = 0.0

                    spend_record = {
                        'po': po_val,
                        'tbm': tbm_val,
                        'product': prod_val,
                        'crop': crop_val,
                        'activity': act_val,
                        'num_activities': num_act_val,
                        'total_amount': amt_val,
                        'area': '',
                        'zdgm': ''
                    }

                    if po_val not in tbm_data:
                        tbm_data[po_val] = []
                    tbm_data[po_val].append(spend_record)

                    items_to_mark.append({
                        'file_path': ef,
                        'sheet_name': target_sname,
                        'header_row': header_row,
                        'data_rows': [r],
                        'po': po_val,
                        'tbm': tbm_val,
                        'activity': act_val,
                        'product': prod_val,
                        'crop': crop_val,
                        'is_summary_sheet': True
                    })

        if file_had_any_data and not file_has_unmarked:
            skipped_files_count += 1

        wb.close()

    return tbm_data, items_to_mark, skipped_files_count, total_files_scanned


def mark_tbms_as_synced(items_to_mark):
    """
    Marks the processed tables/rows in source All-TBMs summary workbooks as DONE:
    - Writes 'Status' in Column R (Col 18) at header_row
    - Writes 'DONE' in Column R (Col 18) for each data row
    - If 'TBM Amount Summary' sheet exists, also marks Col 8 with 'DONE'
    """
    if not items_to_mark:
        return 0, []

    # Group by file_path
    by_file = {}
    for item in items_to_mark:
        fpath = Path(item['file_path']).resolve()
        if fpath not in by_file:
            by_file[fpath] = []
        by_file[fpath].append(item)

    marked_files_count = 0
    mark_errors = []

    for fpath, file_items in by_file.items():
        if not fpath.exists() or fpath.suffix.lower() not in ['.xlsx', '.xlsm']:
            continue

        try:
            wb = openpyxl.load_workbook(fpath)
        except PermissionError:
            mark_errors.append(f"Cannot mark {fpath.name} as DONE because it is currently open in Excel.")
            continue
        except Exception as ex:
            mark_errors.append(f"Error opening {fpath.name} to mark DONE: {ex}")
            continue

        for item in file_items:
            sname = item['sheet_name']
            if sname not in wb.sheetnames:
                continue

            ws = wb[sname]
            if item.get('is_summary_sheet'):
                # Mark Column H (Col 8) in summary sheet
                ws.column_dimensions['H'].width = 12
                hr = item.get('header_row', 2)
                hdr_cell = ws.cell(hr, 8, "Status")
                hdr_cell.font = STATUS_HDR_FONT
                hdr_cell.alignment = CENTER_ALIGN
                hdr_cell.border = STATUS_BORDER

                for r in item.get('data_rows', []):
                    c = ws.cell(r, 8, "DONE")
                    c.font = STATUS_VAL_FONT
                    c.alignment = CENTER_ALIGN
                    c.border = STATUS_BORDER
            else:
                # Mark Column R (Col 18) in activity sheet
                ws.column_dimensions['R'].width = 12
                hr = item.get('header_row')
                if hr:
                    hdr_cell = ws.cell(hr, 18, "Status")
                    hdr_cell.font = STATUS_HDR_FONT
                    hdr_cell.alignment = CENTER_ALIGN
                    hdr_cell.border = STATUS_BORDER

                for r in item.get('data_rows', []):
                    c = ws.cell(r, 18, "DONE")
                    c.font = STATUS_VAL_FONT
                    c.alignment = CENTER_ALIGN
                    c.border = STATUS_BORDER

        # Also check if 'TBM Amount Summary' sheet is present to mark matching rows in Col 8
        summary_sheet_name = None
        for s in wb.sheetnames:
            if 'tbm amount summary' in s.lower() or 'amount summary' in s.lower():
                summary_sheet_name = s
                break

        if summary_sheet_name:
            ws_sum = wb[summary_sheet_name]
            ws_sum.column_dimensions['H'].width = 12

            # Find header row
            sum_hr = 2
            for r in range(1, min(10, ws_sum.max_row + 1)):
                c1 = str(ws_sum.cell(r, 1).value or '').strip().lower()
                if 'po' in c1:
                    sum_hr = r
                    break

            hdr_c = ws_sum.cell(sum_hr, 8, "Status")
            hdr_c.font = STATUS_HDR_FONT
            hdr_c.alignment = CENTER_ALIGN
            hdr_c.border = STATUS_BORDER

            # Match items
            for item in file_items:
                po_match = item.get('po', '').strip().upper()
                tbm_match = item.get('tbm', '').strip().lower()
                act_norm = normalize_activity(item.get('activity', ''))

                for r in range(sum_hr + 1, ws_sum.max_row + 1):
                    r_po = str(ws_sum.cell(r, 1).value or '').strip().upper()
                    if r_po == po_match:
                        r_tbm = str(ws_sum.cell(r, 2).value or '').strip().lower()
                        r_act = normalize_activity(ws_sum.cell(r, 5).value)
                        if tbm_match in r_tbm or r_tbm in tbm_match:
                            if act_norm == r_act:
                                c = ws_sum.cell(r, 8, "DONE")
                                c.font = STATUS_VAL_FONT
                                c.alignment = CENTER_ALIGN
                                c.border = STATUS_BORDER

        try:
            wb.save(fpath)
            marked_files_count += 1
        except PermissionError:
            mark_errors.append(f"Could not save DONE status to {fpath.name} because it is open in Excel.")
        except Exception as ex:
            mark_errors.append(f"Error saving {fpath.name}: {ex}")
        finally:
            wb.close()

    return marked_files_count, mark_errors
