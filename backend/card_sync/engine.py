"""
Card Sync Engine Orchestrator Module
Sri Balaji Traders Automation System
"""

from pathlib import Path
import openpyxl

try:
    from excel_parser import load_any_workbook
except ImportError:
    import sys
    sys.path.append(str(Path(__file__).resolve().parent.parent))
    from excel_parser import load_any_workbook

from .extractor import extract_all_tbms_spent, mark_tbms_as_synced
from .corteva_sync import sync_corteva_cards_workbook
from .fmc_sync import sync_fmc_cards_workbook


def sync_tbm_with_cards(cards_excel_path, tbm_summary_excel_path, output_path=None, service_charge_percent=5.0, force=False):
    """
    Main entry point to synchronize PO Cards summary with All-TBMs Summary files.
    Accepts either a single All-TBMs summary Excel or a folder containing multiple summaries.
    Preserves existing entries in PO cards (non-destructive append) and marks synced tables/rows as DONE.
    """
    cards_file = Path(cards_excel_path).resolve()
    tbm_path = Path(tbm_summary_excel_path).resolve()

    if not cards_file.exists():
        raise FileNotFoundError(f"PO Cards Summary file not found at: {cards_excel_path}")
    if not tbm_path.exists():
        raise FileNotFoundError(f"TBM Summary file or folder not found at: {tbm_summary_excel_path}")

    # 1. Extract un-synced spent records from All-TBMs summary file(s)
    tbm_data, items_to_mark, skipped_files_count, total_files_scanned = extract_all_tbms_spent(
        tbm_path, force=force
    )

    if not tbm_data:
        if skipped_files_count > 0:
            return {
                "success": True,
                "message": f"All {skipped_files_count} All-TBMs summary file(s) are already marked as DONE. No new spendings to sync.",
                "updatedCards": 0,
                "totalTbmPOs": 0,
                "syncedRecords": 0,
                "markedFilesCount": 0,
                "skippedFilesCount": skipped_files_count,
                "outputPath": str(cards_file),
                "serviceChargePercent": service_charge_percent
            }
        return {
            "success": False,
            "message": f"No valid spending records found in: {tbm_path.name}",
            "updatedCards": 0,
            "totalTbmPOs": 0,
            "syncedRecords": 0,
            "outputPath": str(cards_file)
        }

    # 2. Load Cards Workbook
    wb_cards = load_any_workbook(cards_file)

    # Determine if FMC (stacked 19-row cards) or Corteva (product sheets)
    is_fmc_structure = False
    for sname in wb_cards.sheetnames:
        if sname != 'Sheet1':
            ws = wb_cards[sname]
            if ws.cell(1, 1).value and str(ws.cell(1, 1).value).strip().startswith('500'):
                is_fmc_structure = True
                break

    if is_fmc_structure:
        updated_count = sync_fmc_cards_workbook(wb_cards, tbm_data, service_charge_percent)
    else:
        updated_count = sync_corteva_cards_workbook(wb_cards, tbm_data, service_charge_percent)

    save_target = Path(output_path).resolve() if output_path else cards_file
    try:
        wb_cards.save(save_target)
    except PermissionError:
        wb_cards.close()
        raise PermissionError(
            f"Cannot save PO Cards file '{save_target.name}' because it is open in Microsoft Excel. "
            f"Please close the file and try again."
        )
    finally:
        wb_cards.close()

    # 3. Mark processed tables in All-TBMs summary workbooks as DONE
    marked_files_count, mark_errors = mark_tbms_as_synced(items_to_mark)

    total_records = sum(len(item.get('data_rows', [])) for item in items_to_mark)

    msg = (
        f"Successfully synchronized {updated_count} PO card(s) with {total_records} activity record(s) "
        f"from {marked_files_count} All-TBMs summary file(s) at {service_charge_percent}% service charge!"
    )
    if skipped_files_count > 0:
        msg += f" ({skipped_files_count} file(s) already marked as DONE were skipped)"
    if mark_errors:
        msg += f" Note: {'; '.join(mark_errors)}"

    return {
        "success": True,
        "message": msg,
        "updatedCards": updated_count,
        "totalTbmPOs": len(tbm_data),
        "syncedRecords": total_records,
        "markedFilesCount": marked_files_count,
        "skippedFilesCount": skipped_files_count,
        "outputPath": str(save_target),
        "serviceChargePercent": service_charge_percent
    }
