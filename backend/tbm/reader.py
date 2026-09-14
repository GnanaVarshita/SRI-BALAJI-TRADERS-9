"""
TBM Activity Readers
Extracts activity records from OpenPyXL worksheets, legacy xlrd sheets, and slip formats.
"""

import re
from openpyxl.utils import column_index_from_string

try:
    from common.formatters import clean_str, parse_num, parse_date_intelligent, format_date_xlrd
except ImportError:
    from backend.common.formatters import clean_str, parse_num, parse_date_intelligent, format_date_xlrd

try:
    from .constants import FIELD_KEYWORDS
except ImportError:
    from backend.tbm.constants import FIELD_KEYWORDS


def normalize_header(val):
    """Safely normalizes header cell value by collapsing newlines and extra spaces."""
    if val is None:
        return ""
    s = re.sub(r'[\r\n\t]+', ' ', str(val)).strip()
    return re.sub(r'\s+', ' ', s).lower()


def evaluate_cell_value(ws, val, r=None, c=None, visited=None):
    """
    Recursively evaluates cell values, resolving numbers, numeric strings, and Excel formulas.
    Supports coordinates (e.g. K4), range functions (e.g. SUM(L4:O4), SUBTOTAL(9, ...)),
    and arithmetic (+, -, *, /) with cycle detection.
    """
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return 0.0
    if not s.startswith('='):
        try:
            return float(s.replace(',', ''))
        except ValueError:
            return 0.0

    if visited is None:
        visited = set()
    if r is not None and c is not None:
        if (r, c) in visited:
            return 0.0
        visited.add((r, c))

    # Remove Excel absolute reference dollar signs ($) and convert to uppercase
    formula = s[1:].replace('$', '').strip().upper()

    def repl_sum(m):
        args = m.group(1).split(',')
        tot = 0.0
        for arg in args:
            arg = arg.strip()
            if ':' in arg:
                parts = arg.split(':')
                m1 = re.match(r'([A-Z]+)(\d+)', parts[0].strip())
                m2 = re.match(r'([A-Z]+)(\d+)', parts[1].strip())
                if m1 and m2:
                    c1, r1 = column_index_from_string(m1.group(1)), int(m1.group(2))
                    c2, r2 = column_index_from_string(m2.group(1)), int(m2.group(2))
                    for row in range(min(r1, r2), max(r1, r2) + 1):
                        for col in range(min(c1, c2), max(c1, c2) + 1):
                            tot += evaluate_cell_value(ws, ws.cell(row, col).value, row, col, visited.copy())
            else:
                tot += evaluate_cell_value(ws, arg, None, None, visited.copy())
        return str(tot)

    formula = re.sub(r'SUM\s*\(([^)]+)\)', repl_sum, formula)
    formula = re.sub(r'SUBTOTAL\s*\(\s*9\s*,\s*([^)]+)\)', repl_sum, formula)

    def repl_cell(m):
        c_ref, r_ref = m.group(1), int(m.group(2))
        try:
            c_idx = column_index_from_string(c_ref)
        except ValueError:
            return '0'
        if (r_ref, c_idx) in visited:
            return '0'
        sub_vis = visited.copy()
        sub_vis.add((r_ref, c_idx))
        return str(evaluate_cell_value(ws, ws.cell(r_ref, c_idx).value, r_ref, c_idx, sub_vis))

    formula = re.sub(r'\b([A-Z]+)(\d+)\b', repl_cell, formula)

    cleaned = re.sub(r'[^0-9\.\+\-\*\/\(\)\s]', '', formula)
    if not cleaned.strip():
        return 0.0
    try:
        return float(eval(cleaned, {'__builtins__': None}, {}))
    except Exception:
        return 0.0


def parse_slip_format(ws, default_tbm_name="", default_territory=""):
    """Parses single-card / slip format where keys are in Col 2 and values in Col 4."""
    kv = {}
    for r in range(1, 25):
        k = normalize_header(ws.cell(r, 2).value)
        v = ws.cell(r, 4).value
        if k:
            if 'date' in k: kv['date'] = parse_date_intelligent(v)
            elif 'product' in k: kv['product'] = clean_str(v)
            elif 'crop' in k: kv['crop'] = clean_str(v)
            elif 'activity' in k: kv['activity'] = clean_str(v)
            elif 'place' in k or 'territory' in k: kv['territory'] = clean_str(v)
            elif 'farmers' in k: kv['farmers'] = int(parse_num(v))
            elif 'food' in k: kv['food'] = parse_num(v)
            elif 'chairs' in k or 'tent' in k or 'supplier' in k or 'suppl' in k: kv['tent'] = parse_num(v)
            elif 'auto' in k or 'transport' in k: kv['transport'] = parse_num(v)
            elif 'others' in k or 'gift' in k: kv['others'] = parse_num(v)
            elif 'tbm' in k: kv['tbm'] = clean_str(v)
            elif 'po' in k: kv['po_number'] = clean_str(v)

    if kv.get('activity') or kv.get('product') or kv.get('date'):
        tbm_val = kv.get('tbm') or default_tbm_name
        terr_val = kv.get('territory') or default_territory
        tent = kv.get('tent', 0.0)
        food = kv.get('food', 0.0)
        transport = kv.get('transport', 0.0)
        others = kv.get('others', 0.0)
        total = tent + food + transport + others
        return {
            'date': kv.get('date', ''),
            'zdgm': '',
            'tbm': tbm_val,
            'mdo': '',
            'territory': terr_val,
            'product': kv.get('product', ''),
            'crop': kv.get('crop', ''),
            'activity': kv.get('activity', ''),
            'village': '',
            'farmers': kv.get('farmers', 0),
            'tent': tent,
            'food': food,
            'transport': transport,
            'others': others,
            'total': total,
            'po_number': kv.get('po_number', '')
        }
    return None


def extract_activities_from_sheet(ws, default_tbm_name="", default_territory="", file_name=None, skip_done=False):
    """
    Scans OpenPyXL worksheet for all tables / header rows and extracts activity records.
    If skip_done=True, skips rows or tables that have 'DONE' in the Status column (Column R / 18).
    """
    all_activities = []
    header_rows = []
    max_r = ws.max_row or 100
    max_c = min(ws.max_column or 35, 45)

    # 1. Detect all header rows across the sheet
    for r in range(1, max_r + 1):
        col_map = {}
        score = 0
        is_banner = False

        for c in range(1, max_c + 1):
            txt = normalize_header(ws.cell(r, c).value)
            if 'activities expenses by' in txt or 'marketing activities' in txt or 'activities expenses' in txt:
                is_banner = True
                break

        if is_banner:
            continue

        for c in range(1, max_c + 1):
            curr_txt = normalize_header(ws.cell(r, c).value)
            if not curr_txt:
                continue

            # Prioritize sl_no for column 1 or sl/sno labels
            if any(kw == curr_txt or kw in curr_txt for kw in FIELD_KEYWORDS['sl_no']):
                if 'sl_no' not in col_map:
                    col_map['sl_no'] = c
                    score += 1
                    continue

            # Prioritize farmers before mdo to avoid any 'fa' substring collision
            if 'farmers' not in col_map:
                if any(kw == curr_txt or kw in curr_txt for kw in FIELD_KEYWORDS['farmers']):
                    col_map['farmers'] = c
                    score += 1
                    continue

            # Prioritize total to avoid matching 'amount' as zdgm/tbm
            if 'total' not in col_map:
                if any(kw == curr_txt or (len(kw) > 3 and kw in curr_txt) for kw in FIELD_KEYWORDS['total']):
                    col_map['total'] = c
                    score += 1
                    continue

            for field, kw_list in FIELD_KEYWORDS.items():
                if field in ['sl_no', 'farmers', 'total']:
                    continue
                if field not in col_map:
                    matched = False
                    for kw in kw_list:
                        if kw == curr_txt:
                            matched = True
                            break
                        if len(kw) <= 3:
                            if re.search(rf'(?<![a-z0-9]){re.escape(kw)}(?![a-z0-9])', curr_txt):
                                matched = True
                                break
                        else:
                            if kw in curr_txt:
                                matched = True
                                break
                    if matched:
                        col_map[field] = c
                        score += 1
                        break

        # Fallback keyword checks for any unmatched column that contains expense indicators
        if score >= 3:
            for c in range(1, max_c + 1):
                if c in col_map.values():
                    continue
                curr_txt = normalize_header(ws.cell(r, c).value)
                if not curr_txt:
                    continue
                if 'tent' not in col_map and any(w in curr_txt for w in ['supplier', 'suppl', 'tent', 'hall', 'chair']):
                    col_map['tent'] = c
                elif 'food' not in col_map and any(w in curr_txt for w in ['food', 'snack', 'tiffin', 'meal']):
                    col_map['food'] = c
                elif 'transport' not in col_map and any(w in curr_txt for w in ['transport', 'travel', 'auto']):
                    col_map['transport'] = c
                elif 'others' not in col_map and any(w in curr_txt for w in ['other', 'gift']):
                    col_map['others'] = c

        if score >= 3 and any(k in col_map for k in ['date', 'product', 'activity', 'mdo', 'zdgm', 'territory', 'sl_no', 'farmers']):
            header_rows.append((r, col_map))

    if not header_rows:
        slip_act = parse_slip_format(ws, default_tbm_name, default_territory)
        if slip_act:
            if file_name:
                slip_act['source_file'] = file_name
            return [slip_act]
        return []

    # 2. Extract data rows for each detected table header
    for i, (hr, cmap) in enumerate(header_rows):
        next_hr = header_rows[i + 1][0] if i + 1 < len(header_rows) else max_r + 1

        # Identify boundary columns: farmers and total
        c_farmers = cmap.get('farmers')
        if not c_farmers:
            for c in range(1, max_c + 1):
                txt = normalize_header(ws.cell(hr, c).value)
                if any(kw == txt or kw in txt for kw in FIELD_KEYWORDS['farmers']):
                    c_farmers = c
                    break

        c_total = cmap.get('total')
        if not c_total or (c_farmers and c_total <= c_farmers):
            start_search = (c_farmers + 1) if c_farmers else 1
            for c in range(start_search, max_c + 1):
                txt = normalize_header(ws.cell(hr, c).value)
                if any(kw == txt or (len(kw) > 3 and kw in txt) for kw in FIELD_KEYWORDS['total']) and 'sub' not in txt:
                    c_total = c
                    break

        # Map all activity expense columns strictly between c_farmers and c_total
        expense_col_mapping = []  # list of (col_idx, category)
        if c_farmers and c_total and c_farmers < c_total:
            for c in range(c_farmers + 1, c_total):
                txt = normalize_header(ws.cell(hr, c).value)
                cat = 'others'
                if any(w in txt for w in ['supplier', 'suppl', 'tent', 'hall', 'chair', 'sound', 'audio', 'mic', 'projector', 'stage']):
                    cat = 'tent'
                elif any(w in txt for w in ['food', 'snack', 'tiffin', 'meal', 'refreshment', 'lunch', 'dinner']) or txt in ['expenses', 'expense']:
                    cat = 'food'
                elif any(w in txt for w in ['transport', 'travel', 'auto', 'vehicle', 'cab', 'conveyance', 'diesel', 'petrol', 'trans port', 'trans-port']):
                    cat = 'transport'
                elif any(w in txt for w in ['other', 'gift', 'stationary', 'misc', 'printing']):
                    cat = 'others'
                else:
                    cat = 'others'
                expense_col_mapping.append((c, cat))
        else:
            # Fallback if boundary columns are not found
            for cat in ['tent', 'food', 'transport', 'others']:
                if cat in cmap:
                    expense_col_mapping.append((cmap[cat], cat))

        # Identify status column if present
        c_status = cmap.get('status')
        if not c_status:
            for c in range(16, max_c + 1):
                txt = normalize_header(ws.cell(hr, c).value)
                if txt in ['status', 'state']:
                    c_status = c
                    break
        if not c_status:
            c_status = 18

        for r in range(hr + 1, next_hr):
            if skip_done:
                val_status = str(ws.cell(r, c_status).value or "").strip().upper()
                if val_status == "DONE":
                    continue

            row_vals = [ws.cell(r, c).value for c in range(1, max_c + 1)]
            if not any(row_vals):
                continue

            row_str = ' '.join(clean_str(v).lower() for v in row_vals)
            has_data_content = any(
                clean_str(ws.cell(r, cmap.get(k, 0)).value)
                for k in ['product', 'activity', 'date', 'village']
                if k in cmap and cmap.get(k, 0) <= max_c
            )
            if 'total' in row_str and not has_data_content:
                break

            raw_date = ws.cell(r, cmap.get('date', 0)).value if 'date' in cmap else None
            date_val = parse_date_intelligent(raw_date)

            prod_val = clean_str(ws.cell(r, cmap.get('product', 0)).value) if 'product' in cmap else ''
            act_val = clean_str(ws.cell(r, cmap.get('activity', 0)).value) if 'activity' in cmap else ''
            crop_val = clean_str(ws.cell(r, cmap.get('crop', 0)).value) if 'crop' in cmap else ''
            vlg_val = clean_str(ws.cell(r, cmap.get('village', 0)).value) if 'village' in cmap else ''
            zdgm_val = clean_str(ws.cell(r, cmap.get('zdgm', 0)).value) if 'zdgm' in cmap else ''
            mdo_val = clean_str(ws.cell(r, cmap.get('mdo', 0)).value) if 'mdo' in cmap else ''
            terr_val = clean_str(ws.cell(r, cmap.get('territory', 0)).value) if 'territory' in cmap else ''
            if not terr_val:
                terr_val = default_territory

            tbm_val = clean_str(ws.cell(r, cmap.get('tbm', 0)).value) if 'tbm' in cmap else ''
            if not tbm_val or tbm_val.isdigit():
                tbm_val = default_tbm_name

            farm_val = evaluate_cell_value(ws, ws.cell(r, c_farmers).value, r, c_farmers) if c_farmers else 0.0

            cat_vals = {'tent': 0.0, 'food': 0.0, 'transport': 0.0, 'others': 0.0}
            for col_idx, cat in expense_col_mapping:
                cat_vals[cat] += evaluate_cell_value(ws, ws.cell(r, col_idx).value, r, col_idx)

            tent_val = cat_vals['tent']
            food_val = cat_vals['food']
            trans_val = cat_vals['transport']
            oth_val = cat_vals['others']
            sum_exp = tent_val + food_val + trans_val + oth_val

            tot_val = 0.0
            if c_total and c_total <= max_c:
                tot_val = evaluate_cell_value(ws, ws.cell(r, c_total).value, r, c_total)
            elif 'total' in cmap:
                tot_val = evaluate_cell_value(ws, ws.cell(r, cmap['total']).value, r, cmap['total'])

            # Ensure sum of activities expenses matches tot_val if tot_val was explicitly provided
            if tot_val > 0 and sum_exp < tot_val - 0.01:
                diff = round(tot_val - sum_exp, 2)
                matched_cats = {cat for _, cat in expense_col_mapping}
                if tent_val == 0.0 and 'tent' not in matched_cats:
                    tent_val = diff
                else:
                    oth_val += diff
                sum_exp = tent_val + food_val + trans_val + oth_val

            if tot_val == 0.0 or tot_val < sum_exp - 0.01:
                tot_val = sum_exp

            if not prod_val and not act_val and not date_val and not vlg_val and not crop_val and tot_val == 0.0:
                continue
            if not prod_val and not act_val and not date_val and not vlg_val and not crop_val and not (zdgm_val and terr_val):
                continue

            po_val = clean_str(ws.cell(r, cmap.get('po_number', 0)).value) if 'po_number' in cmap else ''
            if not po_val:
                for c in range(1, max_c + 1):
                    v = clean_str(ws.cell(r, c).value)
                    m = re.search(r'[45]\d{2}[A-Z0-9]{7,20}', v, re.I)
                    if m:
                        po_val = m.group(0).upper()
                        break

            rec = {
                'date': date_val,
                'zdgm': zdgm_val,
                'tbm': tbm_val,
                'mdo': mdo_val,
                'territory': terr_val,
                'product': prod_val,
                'crop': crop_val,
                'activity': act_val,
                'village': vlg_val,
                'farmers': int(farm_val),
                'tent': tent_val,
                'food': food_val,
                'transport': trans_val,
                'others': oth_val,
                'total': tot_val,
                'po_number': po_val,
                '_row_idx': r,
                '_hr_idx': hr
            }
            if file_name:
                rec['source_file'] = file_name
            all_activities.append(rec)

    return all_activities


def extract_activities_from_xlrd_sheet(sh, datemode, tbm_folder_name, file_name):
    """
    Extracts activity records from legacy .xls sheets using xlrd.
    """
    best_row = None
    best_map = {}
    best_score = 0

    max_r = min(sh.nrows, 25)
    for r in range(max_r):
        curr_map = {}
        score = 0
        for c in range(sh.ncols):
            top_txt = normalize_header(sh.cell_value(r - 1, c)) if r > 0 else ""
            curr_txt = normalize_header(sh.cell_value(r, c))
            comb_txt = f"{top_txt} {curr_txt}".strip()
            if not comb_txt:
                continue

            # Prioritize sl_no
            if any(kw == comb_txt or kw in comb_txt for kw in FIELD_KEYWORDS['sl_no']):
                if 'sl_no' not in curr_map:
                    curr_map['sl_no'] = c
                    score += 1
                    continue

            # Prioritize farmers
            if 'farmers' not in curr_map:
                if any(kw == comb_txt or kw in comb_txt for kw in FIELD_KEYWORDS['farmers']):
                    curr_map['farmers'] = c
                    score += 1
                    continue

            # Prioritize total
            if 'total' not in curr_map:
                if any(kw == comb_txt or (len(kw) > 3 and kw in comb_txt) for kw in FIELD_KEYWORDS['total']):
                    curr_map['total'] = c
                    score += 1
                    continue

            for field, kw_list in FIELD_KEYWORDS.items():
                if field in ['sl_no', 'farmers', 'total']:
                    continue
                if field not in curr_map:
                    matched = False
                    for kw in kw_list:
                        if kw == comb_txt:
                            matched = True
                            break
                        if len(kw) <= 3:
                            if re.search(rf'(?<![a-z0-9]){re.escape(kw)}(?![a-z0-9])', comb_txt):
                                matched = True
                                break
                        else:
                            if kw in comb_txt:
                                matched = True
                                break
                    if matched:
                        curr_map[field] = c
                        score += 1
                        break

        # Fallback for expense columns
        if score >= 3:
            for c in range(sh.ncols):
                if c in curr_map.values():
                    continue
                top_txt = normalize_header(sh.cell_value(r - 1, c)) if r > 0 else ""
                curr_txt = normalize_header(sh.cell_value(r, c))
                comb_txt = f"{top_txt} {curr_txt}".strip()
                if not comb_txt:
                    continue
                if 'tent' not in curr_map and any(w in comb_txt for w in ['supplier', 'suppl', 'tent', 'hall', 'chair']):
                    curr_map['tent'] = c
                elif 'food' not in curr_map and any(w in comb_txt for w in ['food', 'snack', 'tiffin', 'meal']):
                    curr_map['food'] = c
                elif 'transport' not in curr_map and any(w in comb_txt for w in ['transport', 'travel', 'auto']):
                    curr_map['transport'] = c
                elif 'others' not in curr_map and any(w in comb_txt for w in ['other', 'gift']):
                    curr_map['others'] = c
        if score > best_score and score >= 3:
            best_score = score
            best_row = r
            best_map = curr_map

    if best_row is None:
        return []

    c_farmers = best_map.get('farmers')
    c_total = best_map.get('total')
    expense_col_mapping = []
    if c_farmers is not None and c_total is not None and c_farmers < c_total:
        for c in range(c_farmers + 1, c_total):
            top_txt = normalize_header(sh.cell_value(best_row - 1, c)) if best_row > 0 else ""
            curr_txt = normalize_header(sh.cell_value(best_row, c))
            comb_txt = f"{top_txt} {curr_txt}".strip()
            cat = 'others'
            if any(w in comb_txt for w in ['supplier', 'suppl', 'tent', 'hall', 'chair', 'sound', 'audio', 'mic', 'projector', 'stage']):
                cat = 'tent'
            elif any(w in comb_txt for w in ['food', 'snack', 'tiffin', 'meal', 'refreshment', 'lunch', 'dinner']) or comb_txt in ['expenses', 'expense']:
                cat = 'food'
            elif any(w in comb_txt for w in ['transport', 'travel', 'auto', 'vehicle', 'cab', 'conveyance', 'diesel', 'petrol', 'trans port', 'trans-port']):
                cat = 'transport'
            elif any(w in comb_txt for w in ['other', 'gift', 'stationary', 'misc', 'printing']):
                cat = 'others'
            else:
                cat = 'others'
            expense_col_mapping.append((c, cat))
    else:
        for cat in ['tent', 'food', 'transport', 'others']:
            if cat in best_map:
                expense_col_mapping.append((best_map[cat], cat))

    sheet_po_number = ""
    for r in range(min(10, sh.nrows)):
        for c in range(min(20, sh.ncols)):
            val = clean_str(sh.cell_value(r, c))
            m = re.search(r'[45]\d{2}[A-Z0-9]{7,20}', val, re.I)
            if m:
                sheet_po_number = m.group(0).upper()
                break
        if sheet_po_number:
            break

    activities = []
    for r in range(best_row + 1, sh.nrows):
        row_vals = [sh.cell_value(r, c) for c in range(sh.ncols)]
        if not any(row_vals):
            continue

        c1_val = clean_str(row_vals[0]).lower()
        c2_val = clean_str(row_vals[1]).lower() if len(row_vals) > 1 else ""
        date_col_idx = best_map.get('date', 2)
        date_c_val = clean_str(sh.cell_value(r, date_col_idx)).lower() if date_col_idx < len(row_vals) else ""

        if c1_val == 'total' or c2_val == 'total' or date_c_val == 'total':
            break

        date_val = format_date_xlrd(sh.cell_value(r, date_col_idx), datemode) if date_col_idx < len(row_vals) else ""
        
        tbm_idx = best_map.get('tbm', 0)
        tbm_val = clean_str(sh.cell_value(r, tbm_idx)) if tbm_idx < len(row_vals) else ""
        if not tbm_val:
            tbm_val = tbm_folder_name

        zdgm_idx = best_map.get('zdgm', 0)
        zdgm_val = clean_str(sh.cell_value(r, zdgm_idx)) if zdgm_idx < len(row_vals) else ""

        mdo_idx = best_map.get('mdo', 0)
        mdo_val = clean_str(sh.cell_value(r, mdo_idx)) if mdo_idx < len(row_vals) else ""

        terr_idx = best_map.get('territory', 0)
        territory_val = clean_str(sh.cell_value(r, terr_idx)) if terr_idx < len(row_vals) else ""

        prod_idx = best_map.get('product', 0)
        product_val = clean_str(sh.cell_value(r, prod_idx)) if prod_idx < len(row_vals) else ""

        crop_idx = best_map.get('crop', 0)
        crop_val = clean_str(sh.cell_value(r, crop_idx)) if crop_idx < len(row_vals) else ""

        act_idx = best_map.get('activity', 0)
        activity_val = clean_str(sh.cell_value(r, act_idx)) if act_idx < len(row_vals) else ""

        vlg_idx = best_map.get('village', 0)
        village_val = clean_str(sh.cell_value(r, vlg_idx)) if vlg_idx < len(row_vals) else ""

        farm_idx = best_map.get('farmers', 0)
        farmers_val = parse_num(sh.cell_value(r, farm_idx)) if farm_idx < len(row_vals) else 0.0

        cat_vals = {'tent': 0.0, 'food': 0.0, 'transport': 0.0, 'others': 0.0}
        for col_idx, cat in expense_col_mapping:
            if col_idx < len(row_vals):
                cat_vals[cat] += parse_num(sh.cell_value(r, col_idx))

        tent_val = cat_vals['tent']
        food_val = cat_vals['food']
        transport_val = cat_vals['transport']
        others_val = cat_vals['others']
        calc_total = tent_val + food_val + transport_val + others_val

        tot_idx = best_map.get('total')
        total_val = parse_num(sh.cell_value(r, tot_idx)) if tot_idx is not None and tot_idx < len(row_vals) else 0.0
        if total_val == 0.0 or total_val < calc_total - 0.01:
            total_val = calc_total

        po_idx = best_map.get('po_number', 0)
        po_val = clean_str(sh.cell_value(r, po_idx)) if po_idx < len(row_vals) else ""
        if not po_val and sheet_po_number:
            po_val = sheet_po_number
        if not po_val:
            for c in range(len(row_vals)):
                v = clean_str(sh.cell_value(r, c))
                m = re.search(r'[45]\d{2}[A-Z0-9]{7,20}', v, re.I)
                if m:
                    po_val = m.group(0).upper()
                    break
        if not po_val and sheet_po_number:
            po_val = sheet_po_number

        if not activity_val and not product_val and total_val == 0.0 and not date_val:
            continue

        activities.append({
            'date': date_val,
            'zdgm': zdgm_val,
            'tbm': tbm_val,
            'mdo': mdo_val,
            'territory': territory_val,
            'product': product_val,
            'crop': crop_val,
            'activity': activity_val,
            'village': village_val,
            'farmers': int(farmers_val),
            'tent': tent_val,
            'food': food_val,
            'transport': transport_val,
            'others': others_val,
            'total': total_val,
            'po_number': po_val,
            'source_file': file_name
        })

    return activities
