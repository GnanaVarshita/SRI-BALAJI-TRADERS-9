"""
TBM Consolidator
Consolidates activities across multiple TBMs into an All-TBMs Master Summary workbook.
"""

import re
import csv
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    import xlrd
except ImportError:
    xlrd = None

from .reader import extract_activities_from_sheet, extract_activities_from_xlrd_sheet


def create_table_styles():
    """Styles for consolidated summary tables."""
    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    total_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    summary_border = Border(
        left=Side(style='medium', color='702000'),
        right=Side(style='medium', color='702000'),
        top=Side(style='medium', color='702000'),
        bottom=Side(style='double', color='702000')
    )

    brown_color = '702000'

    return {
        'thin_border': thin_border,
        'total_border': total_border,
        'summary_border': summary_border,
        'subtotal_fill': PatternFill(start_color='F2DCDB', end_color='F2DCDB', fill_type='solid'),
        'po_total_fill': PatternFill(start_color='DDD9C4', end_color='DDD9C4', fill_type='solid'),
        'title_font': Font(name='Calibri', size=11, bold=True, color=brown_color),
        'hdr_font': Font(name='Calibri', size=10, bold=True, color=brown_color),
        'brown_text_font': Font(name='Calibri', size=10, color=brown_color),
        'date_font': Font(name='Calibri', size=10, color=brown_color),
        'regular_font': Font(name='Calibri', size=10),
        'bold_font': Font(name='Calibri', size=10, bold=True),
        'po_font': Font(name='Calibri', size=10, bold=True, color=brown_color),
        'subtotal_font': Font(name='Calibri', size=10, bold=True, color=brown_color),
        'po_grand_font': Font(name='Calibri', size=11, bold=True, color=brown_color)
    }


def write_table_to_sheet(ws, start_row, tbm_name, territory, rows_data, po_number, styles):
    """
    Writes a single formatted TBM activity table to worksheet at start_row matching screenshot layout.
    Returns (next_available_row, total_cell_row_idx).
    """
    title_row = start_row
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=16)
    title_cell = ws.cell(title_row, 1)
    terr_str = f"-{territory.upper()}" if territory else ""
    title_cell.value = f"MARKETING ACTIVITIES EXPENSES-{tbm_name.upper()}{terr_str}"
    title_cell.font = styles['title_font']
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[title_row].height = 24

    for c in range(1, 18):
        ws.cell(title_row, c).border = styles['thin_border']

    hdr_row = title_row + 1
    headers = [
        "SI No", "Date", "ZDGM", "TBM", "MDO", "Territory", "Product", "Crop", "Activity", "Village",
        "No. of Farmers", "Tent/Hall/Chairs Suppliers Charges", "Food Expenses", "Transport", "Others/Gifts", "Total", "PO Number"
    ]
    ws.row_dimensions[hdr_row].height = 32
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(hdr_row, col_idx)
        cell.value = h_text
        cell.font = styles['hdr_font']
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = styles['thin_border']

    current_r = hdr_row + 1
    for idx, item in enumerate(rows_data, start=1):
        ws.row_dimensions[current_r].height = 20

        c_sno = ws.cell(current_r, 1, idx)
        c_sno.font = styles['bold_font']
        c_sno.alignment = Alignment(horizontal='center', vertical='center')

        c_date = ws.cell(current_r, 2, item.get('date', ''))
        c_date.font = styles['date_font']
        c_date.alignment = Alignment(horizontal='center', vertical='center')

        for c_idx, key in [(3, 'zdgm'), (4, 'tbm'), (5, 'mdo'), (6, 'territory'), (7, 'product'), (8, 'crop'), (9, 'activity'), (10, 'village')]:
            val = item.get(key, '')
            if key == 'tbm' and not val:
                val = tbm_name
            if key == 'territory' and not val:
                val = territory
            c_txt = ws.cell(current_r, c_idx, val)
            c_txt.font = styles['brown_text_font']
            c_txt.alignment = Alignment(horizontal='left' if c_idx in [3,4,5,10] else 'center', vertical='center')

        c_farmers = ws.cell(current_r, 11, item.get('farmers', 0))
        c_farmers.font = styles['regular_font']
        c_farmers.alignment = Alignment(horizontal='center', vertical='center')
        c_farmers.number_format = '#,##0'

        for c_offset, key in enumerate(['tent', 'food', 'transport', 'others'], start=12):
            amt_val = item.get(key, 0.0)
            c_amt = ws.cell(current_r, c_offset, amt_val if amt_val > 0 else "")
            c_amt.font = styles['regular_font']
            c_amt.alignment = Alignment(horizontal='right', vertical='center')
            if amt_val > 0:
                c_amt.number_format = '#,##0'

        c_tot = ws.cell(current_r, 16, f"=SUM(L{current_r}:O{current_r})")
        c_tot.font = styles['regular_font']
        c_tot.alignment = Alignment(horizontal='right', vertical='center')
        c_tot.number_format = '#,##0'

        row_po = item.get('po_number') or po_number or ''
        c_po = ws.cell(current_r, 17, row_po)
        c_po.font = styles['po_font']
        c_po.alignment = Alignment(horizontal='center', vertical='center')

        for c in range(1, 18):
            ws.cell(current_r, c).border = styles['thin_border']

        current_r += 1

    po_extra_row = current_r
    ws.row_dimensions[po_extra_row].height = 18
    for c in range(1, 18):
        ws.cell(po_extra_row, c).border = styles['thin_border']
    c_po_last = ws.cell(po_extra_row, 17, po_number or "")
    c_po_last.font = styles['po_font']
    c_po_last.alignment = Alignment(horizontal='center', vertical='center')
    current_r += 1

    tot_row = current_r
    ws.row_dimensions[tot_row].height = 22
    for c in range(1, 18):
        ws.cell(tot_row, c).border = styles['total_border']

    lbl_cell = ws.cell(tot_row, 15, "TOTAL")
    lbl_cell.font = styles['bold_font']
    lbl_cell.alignment = Alignment(horizontal='right', vertical='center')

    data_start_r = hdr_row + 1
    data_end_r = po_extra_row - 1

    sum_tot_cell = ws.cell(tot_row, 16, f"=SUM(P{data_start_r}:P{data_end_r})")
    sum_tot_cell.font = styles['bold_font']
    sum_tot_cell.alignment = Alignment(horizontal='right', vertical='center')
    sum_tot_cell.number_format = '#,##0'

    return tot_row + 3, tot_row


def create_tbm_amount_summary_sheet(wb, all_extracted_rows, styles):
    """
    Creates a summary worksheet at the end of the workbook listing total amount used by TBMs,
    organized by: PO, TBM, Product, Crop, Activity, No. of Activities, Total Amount.
    """
    sheet_title = "TBM Amount Summary"
    if sheet_title in wb.sheetnames:
        wb.remove(wb[sheet_title])

    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True

    col_widths = {
        'A': 22,  # PO Number
        'B': 26,  # TBM Name
        'C': 18,  # Product
        'D': 16,  # Crop
        'E': 20,  # Activity
        'F': 18,  # No. of Activities
        'G': 20   # Total Amount
    }
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w

    summary_map = {}
    for r in all_extracted_rows:
        po_num = str(r.get('po_number', '')).strip().upper() or "NO PO"
        tbm_name = str(r.get('tbm', '')).strip().title() or "TBM"
        prod = str(r.get('product', '')).strip().title() or "General Product"
        crop = str(r.get('crop', '')).strip().title() or "General Crop"
        act = str(r.get('activity', '')).strip().upper() or "GENERAL"
        tot_amt = float(r.get('total', 0.0) or 0.0)

        key = (po_num, tbm_name, prod, crop, act)
        if key not in summary_map:
            summary_map[key] = {
                'count': 0,
                'total_amount': 0.0
            }
        summary_map[key]['count'] += 1
        summary_map[key]['total_amount'] += tot_amt

    # Title Header
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    title_cell = ws.cell(1, 1, "SUMMARY OF EXPENSES & ACTIVITIES BY TBM")
    title_cell.font = styles['po_grand_font']
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    for c in range(1, 8):
        ws.cell(1, c).border = styles['thin_border']

    # Table Column Headers
    headers = ["PO", "TBM", "Product", "Crop", "Activity", "No. of Activities", "Total Amount"]
    ws.row_dimensions[2].height = 28
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(2, col_idx, h_text)
        cell.font = styles['hdr_font']
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = styles['thin_border']
        cell.fill = PatternFill(start_color='F2DCDB', end_color='F2DCDB', fill_type='solid')

    current_r = 3
    sorted_keys = sorted(summary_map.keys(), key=lambda k: (0 if k[0] != 'NO PO' else 1, k[0], k[1], k[2], k[4]))

    for key in sorted_keys:
        po_num, tbm_name, prod, crop, act = key
        data = summary_map[key]

        ws.row_dimensions[current_r].height = 20

        c_po = ws.cell(current_r, 1, po_num)
        c_po.font = styles['po_font']
        c_po.alignment = Alignment(horizontal='center', vertical='center')

        c_tbm = ws.cell(current_r, 2, tbm_name)
        c_tbm.font = styles['brown_text_font']
        c_tbm.alignment = Alignment(horizontal='left', vertical='center')

        c_prod = ws.cell(current_r, 3, prod)
        c_prod.font = styles['brown_text_font']
        c_prod.alignment = Alignment(horizontal='center', vertical='center')

        c_crop = ws.cell(current_r, 4, crop)
        c_crop.font = styles['brown_text_font']
        c_crop.alignment = Alignment(horizontal='center', vertical='center')

        c_act = ws.cell(current_r, 5, act)
        c_act.font = styles['brown_text_font']
        c_act.alignment = Alignment(horizontal='center', vertical='center')

        c_cnt = ws.cell(current_r, 6, data['count'])
        c_cnt.font = styles['regular_font']
        c_cnt.alignment = Alignment(horizontal='center', vertical='center')
        c_cnt.number_format = '#,##0'

        c_tot = ws.cell(current_r, 7, data['total_amount'])
        c_tot.font = styles['regular_font']
        c_tot.alignment = Alignment(horizontal='right', vertical='center')
        c_tot.number_format = '#,##0'

        for c in range(1, 8):
            ws.cell(current_r, c).border = styles['thin_border']

        current_r += 1

    grand_r = current_r
    ws.row_dimensions[grand_r].height = 24
    ws.merge_cells(start_row=grand_r, start_column=1, end_row=grand_r, end_column=5)
    
    lbl_g = ws.cell(grand_r, 1, "GRAND TOTAL")
    lbl_g.font = styles['po_grand_font']
    lbl_g.alignment = Alignment(horizontal='right', vertical='center')

    for c in range(1, 8):
        ws.cell(grand_r, c).border = styles['summary_border']
        ws.cell(grand_r, c).fill = styles['po_total_fill']

    c_cnt_tot = ws.cell(grand_r, 6, f"=SUM(F3:F{grand_r-1})")
    c_cnt_tot.font = styles['po_grand_font']
    c_cnt_tot.alignment = Alignment(horizontal='center', vertical='center')
    c_cnt_tot.number_format = '#,##0'

    c_amt_tot = ws.cell(grand_r, 7, f"=SUM(G3:G{grand_r-1})")
    c_amt_tot.font = styles['po_grand_font']
    c_amt_tot.alignment = Alignment(horizontal='right', vertical='center')
    c_amt_tot.number_format = '#,##0'


def generate_tbm_summary(tbm_folder_path, output_path=None, priority_po_list=None):
    """
    Main function to scan TBM summary folder, parse activities (supporting .xlsx and legacy .xls files),
    group by PO, Activity, TBM, format into max 9 tables per sheet, with Activity Subtotals and PO Grand Totals.
    """
    tbm_dir = Path(tbm_folder_path).resolve()
    if not tbm_dir.exists() or not tbm_dir.is_dir():
        raise ValueError(f"TBM Summary folder path does not exist: {tbm_folder_path}")

    if not output_path:
        territory_name = "Nandyala"
        for p in tbm_dir.parts:
            p_up = p.upper()
            if any(k in p_up for k in ['KURNOOL', 'NELLORE', 'NANDYAL', 'NANDYALA', 'SURYAPET']):
                territory_name = p.title()
                break
        if territory_name == "Nandyala" and tbm_dir.name and tbm_dir.name != "TBM s Summary":
            territory_name = tbm_dir.name.replace(" ", "-")

        output_filename = f"{territory_name}-All-TBMs-Summary.xlsx"
        output_path = tbm_dir / output_filename
    else:
        output_path = Path(output_path).resolve()

    priority_pos = set()
    if priority_po_list:
        if isinstance(priority_po_list, str):
            tokens = re.split(r'[\s,\n\r]+', priority_po_list)
        else:
            tokens = priority_po_list
        for tok in tokens:
            cleaned = str(tok).strip().upper()
            if cleaned:
                priority_pos.add(cleaned)

    all_extracted_rows = []
    tbm_subfolders = [d for d in tbm_dir.iterdir() if d.is_dir()]
    if not tbm_subfolders:
        tbm_subfolders = [tbm_dir]

    for subfolder in tbm_subfolders:
        tbm_name = subfolder.name if subfolder != tbm_dir else "TBM"
        excel_files = [
            f for f in subfolder.rglob("*")
            if f.is_file() and (f.suffix.lower() in ['.xlsx', '.xls', '.xlsm', '.xlsb', '.csv', ''] or 'excel' in f.name.lower() or 'report' in f.name.lower() or 'activity' in f.name.lower())
        ]

        for ef in excel_files:
            if ef.resolve() == output_path.resolve():
                continue
            if "-All-TBMs-Summary" in ef.name:
                continue

            suf = ef.suffix.lower()

            if suf in ['.xlsx', '.xlsm']:
                try:
                    wb_in = openpyxl.load_workbook(ef, data_only=True)
                    target_sheets = []
                    if len(wb_in.sheetnames) >= 2:
                        target_sheets = [wb_in.sheetnames[1], wb_in.sheetnames[0]]
                    else:
                        target_sheets = wb_in.sheetnames

                    extracted_for_file = []
                    for sname in target_sheets:
                        ws_in = wb_in[sname]
                        sheet_activities = extract_activities_from_sheet(ws_in, tbm_name, file_name=ef.name)
                        if sheet_activities:
                            extracted_for_file = sheet_activities
                            break
                    
                    all_extracted_rows.extend(extracted_for_file)
                    wb_in.close()
                except Exception as e:
                    print(f"Error reading {suf} file {ef.name}: {e}")

            elif suf == '.xls':
                if xlrd is None:
                    print(f"Skipping .xls file {ef.name} because xlrd is not installed")
                    continue
                try:
                    wb_xls = xlrd.open_workbook(ef)
                    for sname in wb_xls.sheet_names():
                        sh_xls = wb_xls.sheet_by_name(sname)
                        sheet_activities = extract_activities_from_xlrd_sheet(sh_xls, wb_xls.datemode, tbm_name, ef.name)
                        all_extracted_rows.extend(sheet_activities)
                except Exception as e:
                    print(f"Error reading .xls file {ef.name}: {e}")

            elif suf == '.csv':
                try:
                    with open(ef, 'r', encoding='utf-8-sig', errors='replace') as csv_file:
                        reader = list(csv.reader(csv_file))
                        class VirtualCsvSheet:
                            def __init__(self, data):
                                self.nrows = len(data)
                                self.ncols = max((len(r) for r in data), default=0)
                                self.data = data
                            def cell_value(self, r, c):
                                if r < len(self.data) and c < len(self.data[r]):
                                    return self.data[r][c]
                                return ""
                        v_sheet = VirtualCsvSheet(reader)
                        sheet_activities = extract_activities_from_xlrd_sheet(v_sheet, 0, tbm_name, ef.name)
                        all_extracted_rows.extend(sheet_activities)
                except Exception as e:
                    print(f"Error reading .csv file {ef.name}: {e}")

            else:
                parsed = False
                try:
                    wb_in = openpyxl.load_workbook(ef, data_only=True)
                    for sname in wb_in.sheetnames:
                        ws_in = wb_in[sname]
                        sheet_activities = extract_activities_from_sheet(ws_in, tbm_name, file_name=ef.name)
                        all_extracted_rows.extend(sheet_activities)
                    wb_in.close()
                    parsed = True
                except Exception:
                    pass

                if not parsed and xlrd is not None:
                    try:
                        wb_xls = xlrd.open_workbook(ef)
                        for sname in wb_xls.sheet_names():
                            sh_xls = wb_xls.sheet_by_name(sname)
                            sheet_activities = extract_activities_from_xlrd_sheet(sh_xls, wb_xls.datemode, tbm_name, ef.name)
                            all_extracted_rows.extend(sheet_activities)
                    except Exception:
                        pass

    if not all_extracted_rows:
        return {
            "success": False,
            "message": f"No activity records found in Excel files inside {tbm_dir.name}",
            "tables_count": 0,
            "sheets_count": 0
        }

    priority_category = []
    unlisted_category = []
    no_po_category = []

    for row in all_extracted_rows:
        po_num = str(row.get('po_number', '')).strip().upper()
        if not po_num:
            no_po_category.append(row)
        else:
            if priority_pos:
                if po_num in priority_pos:
                    priority_category.append(row)
                else:
                    unlisted_category.append(row)
            else:
                priority_category.append(row)

    def build_hierarchical_groups(rows_list):
        po_dict = {}
        for r in rows_list:
            po_num = str(r.get('po_number', '')).strip().upper() or "NO_PO"
            prod = str(r.get('product', '')).strip().title() or "General Product"
            act = str(r.get('activity', '')).strip().upper() or "GENERAL"
            tbm = str(r.get('tbm', '')).strip().title() or "TBM"

            if po_num not in po_dict:
                po_dict[po_num] = {}
            if act not in po_dict[po_num]:
                po_dict[po_num][act] = {}
            tbm_key = (prod, tbm)
            if tbm_key not in po_dict[po_num][act]:
                po_dict[po_num][act][tbm_key] = []
            po_dict[po_num][act][tbm_key].append(r)
        return po_dict

    styles = create_table_styles()

    if output_path.exists():
        try:
            wb = openpyxl.load_workbook(output_path)
        except Exception:
            wb = openpyxl.Workbook()
            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

    def add_po_hierarchy_to_sheet_series(prefix, rows_list):
        if not rows_list:
            return 0

        po_dict = build_hierarchical_groups(rows_list)
        matching_sheets = [s for s in wb.sheetnames if s.startswith(prefix)]
        
        curr_sheet_idx = 1
        if matching_sheets:
            for ms in matching_sheets:
                m = re.search(r'\d+', ms)
                if m:
                    idx = int(m.group(0))
                    if idx > curr_sheet_idx:
                        curr_sheet_idx = idx

        ws = None
        current_row = 1
        tables_on_current_sheet = 0
        total_tables_created = 0

        def get_or_create_sheet(idx):
            name = f"{prefix}{idx}"
            if name in wb.sheetnames:
                s = wb[name]
                max_used = 1
                for r in range(s.max_row, 0, -1):
                    if any(s.cell(r, c).value for c in range(1, 18)):
                        max_used = r
                        break
                return s, max_used + 4
            else:
                s = wb.create_sheet(title=name)
                s.views.sheetView[0].showGridLines = True
                col_w = {'A': 6, 'B': 12, 'C': 16, 'D': 16, 'E': 14, 'F': 14, 'G': 14, 'H': 12, 'I': 12, 'J': 14, 'K': 12, 'L': 15, 'M': 14, 'N': 12, 'O': 12, 'P': 14, 'Q': 18}
                for c_letter, w in col_w.items():
                    s.column_dimensions[c_letter].width = w
                return s, 1

        ws, current_row = get_or_create_sheet(curr_sheet_idx)

        for po_num, acts_dict in po_dict.items():
            display_po = po_num if po_num != "NO_PO" else ""
            po_subtotal_rows_on_sheet = []

            for act, tbm_groups in acts_dict.items():
                act_table_rows_on_sheet = []
                for (prod, tbm), act_rows in tbm_groups.items():
                    if tables_on_current_sheet >= 9:
                        curr_sheet_idx += 1
                        ws, current_row = get_or_create_sheet(curr_sheet_idx)
                        tables_on_current_sheet = 0
                        act_table_rows_on_sheet = []
                        po_subtotal_rows_on_sheet = []

                    terr = act_rows[0].get('territory', '') if act_rows else ''
                    next_r, tbl_tot_r = write_table_to_sheet(ws, current_row, tbm, terr, act_rows, display_po, styles)
                    act_table_rows_on_sheet.append(tbl_tot_r)
                    current_row = next_r
                    tables_on_current_sheet += 1
                    total_tables_created += 1

                if act_table_rows_on_sheet:
                    subtot_row = current_row
                    ws.row_dimensions[subtot_row].height = 24
                    ws.merge_cells(start_row=subtot_row, start_column=1, end_row=subtot_row, end_column=15)
                    lbl = ws.cell(subtot_row, 1, f"TOTAL EXPENSES FOR ACTIVITY: {act.upper()}")
                    lbl.font = styles['subtotal_font']
                    lbl.fill = styles['subtotal_fill']
                    lbl.alignment = Alignment(horizontal='right', vertical='center')
                    for c in range(1, 18):
                        ws.cell(subtot_row, c).border = styles['thin_border']
                        if c <= 15: ws.cell(subtot_row, c).fill = styles['subtotal_fill']
                    tot_refs = [f"P{r}" for r in act_table_rows_on_sheet]
                    c_tot = ws.cell(subtot_row, 16, f"=SUM({','.join(tot_refs)})")
                    c_tot.font = styles['subtotal_font']
                    c_tot.fill = styles['subtotal_fill']
                    c_tot.alignment = Alignment(horizontal='right', vertical='center')
                    c_tot.number_format = '#,##0'
                    ws.cell(subtot_row, 17, display_po).font = styles['subtotal_font']
                    ws.cell(subtot_row, 17).fill = styles['subtotal_fill']
                    ws.cell(subtot_row, 17).alignment = Alignment(horizontal='center', vertical='center')
                    po_subtotal_rows_on_sheet.append(subtot_row)
                    current_row = subtot_row + 3

            if po_subtotal_rows_on_sheet:
                grand_row = current_row
                ws.row_dimensions[grand_row].height = 26
                ws.merge_cells(start_row=grand_row, start_column=1, end_row=grand_row, end_column=15)
                lbl_g = ws.cell(grand_row, 1, f"GRAND TOTAL FOR PO: {display_po}" if display_po else "GRAND TOTAL FOR UNASSIGNED ACTIVITIES")
                lbl_g.font = styles['po_grand_font']
                lbl_g.fill = styles['po_total_fill']
                lbl_g.alignment = Alignment(horizontal='right', vertical='center')
                for c in range(1, 18):
                    ws.cell(grand_row, c).border = styles['summary_border']
                    if c <= 15: ws.cell(grand_row, c).fill = styles['po_total_fill']
                po_sub_refs = [f"P{r}" for r in po_subtotal_rows_on_sheet]
                c_g = ws.cell(grand_row, 16, f"=SUM({','.join(po_sub_refs)})")
                c_g.font = styles['po_grand_font']
                c_g.fill = styles['po_total_fill']
                c_g.alignment = Alignment(horizontal='right', vertical='center')
                c_g.number_format = '#,##0'
                ws.cell(grand_row, 17, display_po).font = styles['po_grand_font']
                ws.cell(grand_row, 17).fill = styles['po_total_fill']
                ws.cell(grand_row, 17).alignment = Alignment(horizontal='center', vertical='center')
                current_row = grand_row + 4

        return total_tables_created

    n_pri_tbls = add_po_hierarchy_to_sheet_series("Sheet", priority_category)
    n_unl_tbls = add_po_hierarchy_to_sheet_series("Unlisted POs ", unlisted_category)
    n_nopo_tbls = add_po_hierarchy_to_sheet_series("No PO ", no_po_category)

    create_tbm_amount_summary_sheet(wb, all_extracted_rows, styles)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()

    total_tables = n_pri_tbls + n_unl_tbls + n_nopo_tbls
    return {
        "success": True,
        "message": f"Successfully consolidated {len(all_extracted_rows)} activity records across {total_tables} table(s) with Activity Subtotals, PO Grand Totals, and TBM Amount Summary sheet in {output_path.name}!",
        "outputPath": str(output_path),
        "totalActivities": len(all_extracted_rows),
        "totalTables": total_tables,
        "priorityTables": n_pri_tbls,
        "unlistedTables": n_unl_tbls,
        "noPoTables": n_nopo_tbls,
        "sheetsCount": len(wb.sheetnames)
    }
