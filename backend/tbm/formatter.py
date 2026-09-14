"""
TBM Workbook Formatter
Applies standard green-bordered grouped layout with PO & Grand Totals to TBM summaries.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from .constants import STANDARD_TBM_HEADERS, STANDARD_TBM_COL_WIDTHS
from .reader import extract_activities_from_sheet


def create_green_styles():
    """Creates styles matching the standard layout with bold green headings and clean borders."""
    green_color = '006100'  # Dark Excel green
    
    thin_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    total_border = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='006100'),
        bottom=Side(style='double', color='006100')
    )

    return {
        'thin_border': thin_border,
        'total_border': total_border,
        'title_font': Font(name='Calibri', size=11, bold=True, color=green_color),
        'hdr_font': Font(name='Calibri', size=10, bold=True, color=green_color),
        'green_text_font': Font(name='Calibri', size=10, color=green_color),
        'green_bold_font': Font(name='Calibri', size=10, bold=True, color=green_color),
        'total_lbl_font': Font(name='Calibri', size=11, bold=True, color=green_color),
        'total_val_font': Font(name='Calibri', size=11, bold=True, color=green_color)
    }


def build_grouped_tables(activities):
    """
    Groups activities hierarchically:
    1. By PO Number (POs first, then unassigned/NO_PO)
    2. Within PO, by (Product, Activity)
    """
    grouped_map = {}
    for r in activities:
        po_num = str(r.get('po_number', '')).strip().upper()
        prod = str(r.get('product', '')).strip().title() or "General Product"
        act = str(r.get('activity', '')).strip().title() or "General Activity"
        terr = str(r.get('territory', '')).strip().title()
        tbm = str(r.get('tbm', '')).strip().title()

        group_key = (po_num, prod, act)
        if group_key not in grouped_map:
            grouped_map[group_key] = {
                'po_number': po_num,
                'product': prod,
                'activity': act,
                'territory': terr,
                'tbm': tbm,
                'rows': []
            }
        
        if terr and not grouped_map[group_key]['territory']:
            grouped_map[group_key]['territory'] = terr
        if tbm and not grouped_map[group_key]['tbm']:
            grouped_map[group_key]['tbm'] = tbm

        grouped_map[group_key]['rows'].append(r)

    def sort_key(item):
        po, prod, act = item[0]
        is_no_po = 1 if not po else 0
        return (is_no_po, po, prod, act)

    return [v for k, v in sorted(grouped_map.items(), key=sort_key)]


def write_formatted_group_table(ws, start_row, group_data, styles):
    """
    Writes a single formatted table for a (PO, Product, Activity) group.
    Returns (next_start_r, tot_row).
    """
    tbm_name = group_data.get('tbm', '')
    territory = group_data.get('territory', '')
    po_number = group_data.get('po_number', '')
    rows = group_data.get('rows', [])

    title_parts = ["Activities Expenses by"]
    if territory:
        title_parts.append(territory)
    if tbm_name and not tbm_name.isdigit():
        title_parts.append(f"TBM {tbm_name}")
    elif not territory:
        title_parts.append("TBM")
    
    title_text = " ".join(title_parts)

    title_row = start_row
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=17)
    title_cell = ws.cell(title_row, 1, title_text)
    title_cell.font = styles['title_font']
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[title_row].height = 24

    for c in range(1, 18):
        ws.cell(title_row, c).border = styles['thin_border']

    hdr_row = title_row + 1
    ws.row_dimensions[hdr_row].height = 32
    for col_idx, h_text in enumerate(STANDARD_TBM_HEADERS, start=1):
        cell = ws.cell(hdr_row, col_idx, h_text)
        cell.font = styles['hdr_font']
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = styles['thin_border']

    current_r = hdr_row + 1
    data_start_r = current_r

    for idx, item in enumerate(rows, start=1):
        ws.row_dimensions[current_r].height = 20

        c_sno = ws.cell(current_r, 1, idx)
        c_sno.font = styles['green_bold_font']
        c_sno.alignment = Alignment(horizontal='center', vertical='center')

        c_date = ws.cell(current_r, 2, item.get('date', ''))
        c_date.font = styles['green_text_font']
        c_date.alignment = Alignment(horizontal='center', vertical='center')

        field_map = [
            (3, 'zdgm', 'left'),
            (4, 'tbm', 'left'),
            (5, 'mdo', 'left'),
            (6, 'territory', 'left'),
            (7, 'product', 'left'),
            (8, 'crop', 'left'),
            (9, 'activity', 'left'),
            (10, 'village', 'left'),
        ]
        for col_idx, key, align_h in field_map:
            val = item.get(key, '')
            if key == 'tbm' and (not val or val.isdigit()):
                val = tbm_name
            if key == 'territory' and not val:
                val = territory
            c_txt = ws.cell(current_r, col_idx, val)
            c_txt.font = styles['green_text_font']
            c_txt.alignment = Alignment(horizontal=align_h, vertical='center')

        c_farmers = ws.cell(current_r, 11, item.get('farmers', 0))
        c_farmers.font = styles['green_text_font']
        c_farmers.alignment = Alignment(horizontal='center', vertical='center')
        c_farmers.number_format = '#,##0'

        for c_offset, key in enumerate(['tent', 'food', 'transport', 'others'], start=12):
            amt_val = item.get(key, 0.0)
            c_amt = ws.cell(current_r, c_offset, amt_val if amt_val > 0 else "")
            c_amt.font = styles['green_text_font']
            c_amt.alignment = Alignment(horizontal='right', vertical='center')
            if amt_val > 0:
                c_amt.number_format = '#,##0'

        c_tot = ws.cell(current_r, 16, f"=SUM(L{current_r}:O{current_r})")
        c_tot.font = styles['green_text_font']
        c_tot.alignment = Alignment(horizontal='right', vertical='center')
        c_tot.number_format = '#,##0'

        row_po = item.get('po_number') or po_number or ''
        c_po = ws.cell(current_r, 17, row_po)
        c_po.font = styles['green_text_font']
        c_po.alignment = Alignment(horizontal='center', vertical='center')

        for c in range(1, 18):
            ws.cell(current_r, c).border = styles['thin_border']

        current_r += 1

    data_end_r = current_r - 1

    empty_row = current_r
    ws.row_dimensions[empty_row].height = 18
    current_r += 1

    tot_row = current_r
    ws.row_dimensions[tot_row].height = 22

    lbl_cell = ws.cell(tot_row, 15, "Total")
    lbl_cell.font = styles['total_lbl_font']
    lbl_cell.alignment = Alignment(horizontal='right', vertical='center')

    sum_tot_cell = ws.cell(tot_row, 16, f"=SUM(P{data_start_r}:P{data_end_r})")
    sum_tot_cell.font = styles['total_val_font']
    sum_tot_cell.alignment = Alignment(horizontal='right', vertical='center')
    sum_tot_cell.number_format = '#,##0'

    next_start_r = tot_row + 3
    return next_start_r, tot_row


def write_po_and_grand_totals_block(ws, start_row, po_totals, styles):
    """Writes a summary table at the end of Sheet2 listing each PO total and the Grand Total."""
    current_r = start_row
    po_summary_rows = []

    for po_num, tot_rows in po_totals.items():
        ws.row_dimensions[current_r].height = 22
        display_po = po_num if po_num != "NO_PO" else "No PO"

        c_po = ws.cell(current_r, 15, display_po)
        c_po.font = styles['green_bold_font']
        c_po.alignment = Alignment(horizontal='right', vertical='center')
        c_po.border = styles['thin_border']

        refs = [f"P{r}" for r in tot_rows]
        sum_formula = f"=SUM({','.join(refs)})" if len(refs) > 1 else f"={refs[0]}"

        c_tot = ws.cell(current_r, 16, sum_formula)
        c_tot.font = styles['total_val_font']
        c_tot.alignment = Alignment(horizontal='right', vertical='center')
        c_tot.number_format = '#,##0'
        c_tot.border = styles['thin_border']

        po_summary_rows.append(current_r)
        current_r += 1

    ws.row_dimensions[current_r].height = 24
    c_g_lbl = ws.cell(current_r, 15, "Grand Total")
    c_g_lbl.font = styles['total_lbl_font']
    c_g_lbl.alignment = Alignment(horizontal='right', vertical='center')
    c_g_lbl.border = styles['total_border']

    if po_summary_rows:
        g_formula = f"=SUM(P{po_summary_rows[0]}:P{po_summary_rows[-1]})"
    else:
        g_formula = "=0"

    c_g_val = ws.cell(current_r, 16, g_formula)
    c_g_val.font = styles['total_val_font']
    c_g_val.alignment = Alignment(horizontal='right', vertical='center')
    c_g_val.number_format = '#,##0'
    c_g_val.border = styles['total_border']

    return current_r + 2


def is_workbook_already_formatted(wb):
    """
    Checks if a workbook has already been formatted by Step 1 purely based on
    the presence of the 'DONE' status marker in Sheet 1 (Column R / Status column).
    If the user removes or clears 'DONE' from Sheet 1, the workbook is considered
    NOT formatted and will be processed.
    """
    if not wb.sheetnames:
        return False

    ws_raw = wb[wb.sheetnames[0]]
    max_r = min(ws_raw.max_row or 20, 100)
    max_c = min(ws_raw.max_column or 25, 45)

    # Search for "DONE" in Column R (Col 18) or any column from Col 16 onwards
    for r in range(1, max_r + 1):
        for c in range(16, max_c + 1):
            val = str(ws_raw.cell(r, c).value or "").strip().upper()
            if val == "DONE":
                return True

    return False


def mark_sheet1_as_done(ws_raw, styles):
    """
    Marks Sheet 1 as formatted by writing 'Status' in the header row at Column R (Col 18)
    and 'DONE' in each data row of the table.
    """
    thin_border = styles['thin_border']
    status_hdr_font = Font(name='Calibri', size=10, bold=True, color='006100')
    status_val_font = Font(name='Calibri', size=10, bold=True, color='006100')
    center_align = Alignment(horizontal='center', vertical='center')

    max_r = ws_raw.max_row or 50
    max_c = min(ws_raw.max_column or 25, 30)

    # Find table header rows
    header_rows = []
    for r in range(1, min(max_r, 20)):
        row_vals = [str(ws_raw.cell(r, c).value or "").strip().lower() for c in range(1, max_c + 1)]
        row_str = " ".join(row_vals)
        if "activities expenses" in row_str and not any(k in row_str for k in ["sl.no", "sl no", "date", "product"]):
            continue
        if any(k in row_str for k in ["sl.no", "sl no", "s.no", "sno"]) or (
            any(k in row_str for k in ["farmer", "activity"]) and any(k in row_str for k in ["date", "product"])
        ):
            header_rows.append(r)

    if not header_rows:
        header_rows = [2]

    # Column R is column 18
    c_status = 18
    col_letter = openpyxl.utils.get_column_letter(c_status)
    ws_raw.column_dimensions[col_letter].width = 12

    for i, hr in enumerate(header_rows):
        next_hr = header_rows[i + 1] if i + 1 < len(header_rows) else max_r + 1

        # Header cell in Column R
        hdr_cell = ws_raw.cell(hr, c_status, "Status")
        hdr_cell.font = status_hdr_font
        hdr_cell.alignment = center_align
        hdr_cell.border = thin_border

        # Mark all data rows up to Total row or next header
        for r in range(hr + 1, next_hr):
            row_vals = [ws_raw.cell(r, c).value for c in range(1, c_status)]
            if not any(row_vals):
                continue
            row_str = " ".join(str(v).lower() for v in row_vals if v is not None)
            if "total" in row_str and not any(str(ws_raw.cell(r, c).value or "").strip() for c in [2, 7, 8, 9, 10]):
                break

            val_cell = ws_raw.cell(r, c_status, "DONE")
            val_cell.font = status_val_font
            val_cell.alignment = center_align
            val_cell.border = thin_border


def format_tbm_workbook(file_path, default_tbm_name="", default_territory="", force=False):
    """
    Reads Sheet 1 (original raw table), extracts activity records,
    groups them by (PO, Product, Activity), formats them on Sheet 2, and appends
    a PO Totals + Grand Total summary block. Marks Sheet 1 Column R as DONE.
    Skips if already formatted unless force=True.
    """
    file_path = Path(file_path).resolve()
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if not default_tbm_name:
        default_tbm_name = file_path.parent.name if file_path.parent.name != "TBM s Summary" else ""
    if not default_territory:
        for part in file_path.parts:
            p_up = part.upper()
            if any(k in p_up for k in ['KURNOOL', 'NELLORE', 'NANDYAL', 'NANDYALA', 'SURYAPET', 'KAVALI', 'ALLAGADDA', 'ADONI']):
                default_territory = part.replace("-FMC", "").replace(" POs", "").title()
                break

    wb = openpyxl.load_workbook(file_path)

    # Check if already formatted
    if not force and is_workbook_already_formatted(wb):
        wb.close()
        return {
            "success": True,
            "alreadyFormatted": True,
            "skipped": True,
            "message": f"Skipped {file_path.name} (already formatted and marked as DONE)",
            "file": file_path.name,
            "filePath": str(file_path),
            "groupsCount": 0,
            "activitiesCount": 0
        }

    first_sheet_name = wb.sheetnames[0]
    ws_raw = wb[first_sheet_name]
    activities = extract_activities_from_sheet(ws_raw, default_tbm_name, default_territory)

    if not activities and len(wb.sheetnames) > 1:
        for sname in wb.sheetnames[1:]:
            if "formatted" not in sname.lower() and sname != "Sheet2":
                ws_alt = wb[sname]
                alt_acts = extract_activities_from_sheet(ws_alt, default_tbm_name, default_territory)
                if alt_acts:
                    activities.extend(alt_acts)
                    break

    if not activities:
        wb.close()
        return {
            "success": False,
            "message": f"No valid activity records found in {file_path.name}",
            "file": file_path.name,
            "groupsCount": 0,
            "activitiesCount": 0
        }

    groups = build_grouped_tables(activities)

    target_sheet_name = "Sheet2"
    if len(wb.sheetnames) >= 2:
        second_sheet_name = wb.sheetnames[1]
        wb.remove(wb[second_sheet_name])
        ws_target = wb.create_sheet(title=second_sheet_name, index=1)
    else:
        ws_target = wb.create_sheet(title=target_sheet_name, index=1)

    ws_target.views.sheetView[0].showGridLines = True

    for col_let, w in STANDARD_TBM_COL_WIDTHS.items():
        ws_target.column_dimensions[col_let].width = w

    styles = create_green_styles()

    current_r = 1
    po_totals = {}

    for group_data in groups:
        po_num = str(group_data.get('po_number', '')).strip().upper() or "NO_PO"
        current_r, tbl_tot_r = write_formatted_group_table(ws_target, current_r, group_data, styles)
        if po_num not in po_totals:
            po_totals[po_num] = []
        po_totals[po_num].append(tbl_tot_r)

    write_po_and_grand_totals_block(ws_target, current_r, po_totals, styles)

    # Mark Sheet 1 as DONE in Column R
    mark_sheet1_as_done(ws_raw, styles)

    try:
        wb.save(file_path)
    except PermissionError:
        wb.close()
        raise PermissionError(
            f"Cannot save '{file_path.name}' because it is open in Microsoft Excel or another application. "
            f"Please close the file in Excel and try again."
        )
    wb.close()

    return {
        "success": True,
        "alreadyFormatted": False,
        "skipped": False,
        "message": f"Formatted {len(activities)} activities across {len(groups)} table(s) and marked as DONE on {file_path.name}",
        "file": file_path.name,
        "filePath": str(file_path),
        "groupsCount": len(groups),
        "activitiesCount": len(activities),
        "poTotals": list(po_totals.keys())
    }


def format_all_tbm_summaries_in_folder(tbm_folder_path, force=False):
    """Batch formats all TBM Excel summaries in a folder (and its TBM subfolders)."""
    tbm_dir = Path(tbm_folder_path).resolve()
    if not tbm_dir.exists() or not tbm_dir.is_dir():
        raise ValueError(f"TBM Summary folder path does not exist: {tbm_folder_path}")

    excel_files = [
        f for f in tbm_dir.rglob("*")
        if f.is_file() 
        and f.suffix.lower() in ['.xlsx', '.xlsm'] 
        and not f.name.startswith('~$') 
        and "-All-TBMs-Summary" not in f.name
    ]

    if not excel_files:
        return {
            "success": False,
            "message": f"No Excel files found in {tbm_dir.name}",
            "processedFiles": 0,
            "skippedFiles": 0,
            "totalActivities": 0,
            "details": []
        }

    results = []
    total_acts = 0
    total_groups = 0
    success_count = 0
    skipped_count = 0

    for ef in excel_files:
        tbm_folder_name = ef.parent.name if ef.parent != tbm_dir else ""
        try:
            res = format_tbm_workbook(ef, default_tbm_name=tbm_folder_name, force=force)
            results.append(res)
            if res.get("skipped"):
                skipped_count += 1
            elif res.get("success"):
                success_count += 1
                total_acts += res.get("activitiesCount", 0)
                total_groups += res.get("groupsCount", 0)
        except PermissionError as pe:
            results.append({
                "success": False,
                "file": ef.name,
                "message": str(pe)
            })
        except Exception as e:
            results.append({
                "success": False,
                "file": ef.name,
                "message": f"Error formatting {ef.name}: {e}"
            })

    if success_count == 0 and skipped_count > 0:
        summary_msg = f"All {skipped_count} TBM summary file(s) are already formatted and marked as DONE."
    elif success_count > 0 and skipped_count > 0:
        summary_msg = (
            f"Successfully formatted {success_count} file(s) ({total_acts} activities in {total_groups} tables); "
            f"{skipped_count} file(s) were already marked as DONE (skipped)."
        )
    elif success_count > 0 and skipped_count == 0:
        summary_msg = (
            f"Successfully formatted {success_count} / {len(excel_files)} TBM summary file(s) "
            f"({total_acts} activities in {total_groups} tables) and marked them as DONE!"
        )
    else:
        summary_msg = f"Processed {len(excel_files)} file(s), but none could be formatted."

    return {
        "success": (success_count + skipped_count) > 0,
        "message": summary_msg,
        "processedFiles": success_count,
        "skippedFiles": skipped_count,
        "totalFiles": len(excel_files),
        "totalActivities": total_acts,
        "totalTables": total_groups,
        "details": results
    }
