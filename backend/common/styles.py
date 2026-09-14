"""
Excel Styling & Layout Utilities
Sri Balaji Traders Automation System
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def get_base_styles():
    """Returns standard style dictionary used across invoice generations."""
    thin = Side(style='thin', color='A6A6A6')
    black_thin = Side(style='thin', color='000000')
    black_double = Side(style='double', color='000000')
    black_medium = Side(style='medium', color='000000')

    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    box_border = Border(left=black_thin, right=black_thin, top=black_thin, bottom=black_thin)
    total_border = Border(left=black_thin, right=black_thin, top=black_thin, bottom=black_double)
    header_border = Border(left=black_thin, right=black_thin, top=black_medium, bottom=black_medium)

    # Soft SBT background fill
    sbt_fill = PatternFill(start_color='F2EFE9', end_color='F2EFE9', fill_type='solid')

    return {
        'thin_border': thin_border,
        'box_border': box_border,
        'total_border': total_border,
        'header_border': header_border,
        'sbt_fill': sbt_fill,
        'font_title': Font(name='Calibri', size=11, bold=True),
        'font_header': Font(name='Calibri', size=10, bold=True),
        'font_bold': Font(name='Calibri', size=10, bold=True),
        'font_bold_u': Font(name='Calibri', size=10, bold=True, underline='single'),
        'font_regular': Font(name='Calibri', size=10),
        'font_italic': Font(name='Calibri', size=10, italic=True),
        'font_green_title': Font(name='Calibri', size=11, bold=True, color='008000'),
        'font_green_bold': Font(name='Calibri', size=10, bold=True, color='008000'),
        'font_blue_bold': Font(name='Calibri', size=10, bold=True, color='0000FF'),
        'font_sbt_logo': Font(name='Times New Roman', size=26, bold=True),
        'font_sbt_sub': Font(name='Calibri', size=8, bold=True),
        'align_center': Alignment(horizontal='center', vertical='center'),
        'align_center_wrap': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'align_left': Alignment(horizontal='left', vertical='center'),
        'align_right': Alignment(horizontal='right', vertical='center'),
    }


def get_details_styles():
    """Returns standard style dictionary for Details of Bills worksheets."""
    black_thin = Side(style='thin', color='000000')
    black_medium = Side(style='medium', color='000000')
    black_double = Side(style='double', color='000000')

    thin_border = Border(left=black_thin, right=black_thin, top=black_thin, bottom=black_thin)
    box_border = Border(left=black_thin, right=black_thin, top=black_thin, bottom=black_thin)
    summary_border = Border(left=black_thin, right=black_thin, top=black_thin, bottom=black_double)
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    return {
        'thin_border': thin_border,
        'box_border': box_border,
        'summary_border': summary_border,
        'header_fill': header_fill,
        'font_title': Font(name='Calibri', size=11, bold=True),
        'font_summary': Font(name='Calibri', size=10, bold=True),
        'font_header_green': Font(name='Calibri', size=10, bold=True, color='008000'),
        'font_header_red': Font(name='Calibri', size=10, bold=True, color='C00000'),
        'font_header_brown': Font(name='Calibri', size=10, bold=True, color='993300'),
        'font_header_black': Font(name='Calibri', size=10, bold=True, color='000000'),
        'font_bold': Font(name='Calibri', size=10, bold=True),
        'font_regular': Font(name='Calibri', size=10),
        'align_center': Alignment(horizontal='center', vertical='center'),
        'align_center_wrap': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'align_left': Alignment(horizontal='left', vertical='center'),
        'align_right': Alignment(horizontal='right', vertical='center'),
    }


def setup_page_print_fit(ws, print_area=None, orientation="portrait"):
    """
    Configures worksheet page setup for clean A4 printing with balanced side margins.
    """
    if orientation == "landscape":
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    else:
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if print_area:
        ws.print_area = print_area

    # Center horizontally and set balanced side margins for clean spacing
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35


def write_sbt_header_block(ws, start_row, copy_type, styles):
    """
    Writes the standard Sri Balaji Traders top branding block starting at start_row across Columns A to J.
    """
    r1 = start_row
    r2 = start_row + 1
    r3 = start_row + 2
    r4 = start_row + 3
    r5 = start_row + 4
    r6 = start_row + 5
    r7 = start_row + 6
    r8 = start_row + 7

    # Left company info (Cols A-C)
    ws[f'A{r1}'] = "SRI BALAJI TRADERS"
    ws[f'A{r1}'].font = styles['font_title']
    ws[f'A{r2}'] = "#5/387, Gandhi Road, Proddatur,"
    ws[f'A{r2}'].font = styles['font_regular']
    ws[f'A{r3}'] = "YSR District, (Kadapa)"
    ws[f'A{r3}'].font = styles['font_regular']
    ws[f'A{r4}'] = "Andhra Pradesh, 516360"
    ws[f'A{r4}'].font = styles['font_regular']
    ws[f'A{r5}'] = "GSTIN:37BXDPK0359K1ZA"
    ws[f'A{r5}'].font = styles['font_bold']

    # Center SBT logo block (Cols D-G) with soft fill
    ws.merge_cells(f'D{r1}:G{r4}')
    c_logo = ws[f'D{r1}']
    c_logo.value = "SBT"
    c_logo.font = styles['font_sbt_logo']
    c_logo.alignment = styles['align_center']
    
    ws.merge_cells(f'D{r5}:G{r5}')
    c_logosub = ws[f'D{r5}']
    c_logosub.value = "SRI BALAJI TRADERS"
    c_logosub.font = styles['font_sbt_sub']
    c_logosub.alignment = styles['align_center']

    for r_sbt in range(r1, r5 + 1):
        for c_sbt in range(4, 8):
            ws.cell(r_sbt, c_sbt).fill = styles['sbt_fill']

    # Right Contact Info (Cols I-J)
    ws.merge_cells(f'I{r1}:J{r1}')
    ws[f'I{r1}'] = "Prop: K Radha Devi"
    ws[f'I{r1}'].font = styles['font_bold']
    ws[f'I{r1}'].alignment = styles['align_right']

    ws.merge_cells(f'I{r2}:J{r2}')
    ws[f'I{r2}'] = "Cell: 8328588119"
    ws[f'I{r2}'].font = styles['font_bold']
    ws[f'I{r2}'].alignment = styles['align_right']

    ws.merge_cells(f'I{r3}:J{r3}')
    ws[f'I{r3}'] = "9000491388"
    ws[f'I{r3}'].font = styles['font_bold']
    ws[f'I{r3}'].alignment = styles['align_right']

    # Row 5 solid bottom separator across Columns A to J (Cols 1 to 10)
    for col in range(1, 11):
        cell = ws.cell(r5, col)
        cell.border = Border(bottom=Side(style='medium', color='000000'))

    for r in range(r1, r5 + 1):
        ws.row_dimensions[r].height = 15

    # Row 6 spacer
    ws.row_dimensions[r6].height = 6

    # Titles (Row 7): TAX INVOICE & ORIGINAL / DUPLICATE
    ws.row_dimensions[r7].height = 18
    ws.merge_cells(f'E{r7}:G{r7}')
    ws[f'E{r7}'] = "TAX INVOICE"
    ws[f'E{r7}'].font = Font(name='Calibri', size=11, bold=True, underline='single')
    ws[f'E{r7}'].alignment = styles['align_center']

    ws.merge_cells(f'I{r7}:J{r7}')
    ws[f'I{r7}'] = copy_type
    ws[f'I{r7}'].font = Font(name='Calibri', size=11, bold=True, underline='single')
    ws[f'I{r7}'].alignment = styles['align_center']

    # Row 8 spacer
    ws.row_dimensions[r8].height = 6


def write_bank_and_signature_block(ws, start_row, styles):
    """
    Writes the Bank Details and Authorised Signatory block starting directly at start_row across Cols A to J.
    """
    r_bank = start_row
    r_bnk_name = start_row + 1
    r_acc = start_row + 2
    r_ifsc = start_row + 3
    r_branch = start_row + 4
    r_end = start_row + 5

    # Row 1: Bank details:
    ws[f'A{r_bank}'] = "Bank details:"
    ws[f'A{r_bank}'].font = styles['font_bold_u']
    ws.row_dimensions[r_bank].height = 15

    # Row 2: Bank name: Karnataka Bank & For Sri Balaji Traders
    ws[f'A{r_bnk_name}'] = "Bank name: Karnataka Bank"
    ws[f'A{r_bnk_name}'].font = styles['font_bold']
    ws.merge_cells(f'G{r_bnk_name}:J{r_bnk_name}')
    ws[f'G{r_bnk_name}'] = "For Sri Balaji Traders"
    ws[f'G{r_bnk_name}'].font = styles['font_bold']
    ws[f'G{r_bnk_name}'].alignment = styles['align_center']
    ws.row_dimensions[r_bnk_name].height = 15

    # Row 3: A/c No
    ws[f'A{r_acc}'] = "A/c No"
    ws[f'A{r_acc}'].font = styles['font_bold']
    ws.merge_cells(f'B{r_acc}:E{r_acc}')
    ws[f'B{r_acc}'] = ":6187000600001901"
    ws[f'B{r_acc}'].font = styles['font_bold']
    ws.row_dimensions[r_acc].height = 15

    # Row 4: IFSC
    ws[f'A{r_ifsc}'] = "IFSC"
    ws[f'A{r_ifsc}'].font = styles['font_bold']
    ws.merge_cells(f'B{r_ifsc}:E{r_ifsc}')
    ws[f'B{r_ifsc}'] = ":KARB0000618"
    ws[f'B{r_ifsc}'].font = styles['font_bold']
    ws.row_dimensions[r_ifsc].height = 15

    # Row 5: Branch & Authorised Signatory
    ws[f'A{r_branch}'] = "Branch"
    ws[f'A{r_branch}'].font = styles['font_bold']
    ws.merge_cells(f'B{r_branch}:E{r_branch}')
    ws[f'B{r_branch}'] = ":Proddatur"
    ws[f'B{r_branch}'].font = styles['font_bold']
    
    ws.merge_cells(f'G{r_branch}:J{r_branch}')
    ws[f'G{r_branch}'] = "Authorised Signatory"
    ws[f'G{r_branch}'].font = styles['font_bold']
    ws[f'G{r_branch}'].alignment = styles['align_center']
    ws.row_dimensions[r_branch].height = 18

    # Row 6: bottom padding
    ws.row_dimensions[r_end].height = 12

    # Outer border on bank & signature section
    for r in range(r_bank, r_end + 1):
        ws.cell(r, 1).border = Border(left=Side(style='thin', color='000000'))
        ws.cell(r, 10).border = Border(right=Side(style='thin', color='000000'))
    
    # Close bottom border on Row r_end across A:J
    for c in range(1, 11):
        ws.cell(r_end, c).border = Border(
            bottom=Side(style='thin', color='000000'),
            left=Side(style='thin', color='000000') if c == 1 else None,
            right=Side(style='thin', color='000000') if c == 10 else None
        )
    return r_end
