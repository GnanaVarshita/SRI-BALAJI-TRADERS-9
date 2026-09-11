import unittest
import openpyxl
import tempfile
import os
import sys
from pathlib import Path

backend_dir = str(Path(__file__).resolve().parent.parent)
root_dir = str(Path(__file__).resolve().parent.parent.parent)
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)

from backend.tbm.reader import extract_activities_from_sheet
from backend.tbm.formatter import format_tbm_workbook
from backend.tbm.consolidator import generate_tbm_summary


class TestTbmFormatter(unittest.TestCase):
    def test_supplier_charges_recognition(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Setup headers matching DG EXPENSIVES NAIDUPETA.xlsx
        ws.cell(1, 1, "Activities Expenses by Naidupeta Tbm Nagendra")
        headers = [
            'S.NO', 'DATE', 'ZDGM', 'TBM', 'MDO', 'TERRITORY', 'PRODUCT', 'CROP',
            'ACTIVITY', 'VILLAGE', 'NO.OF FARMERS', 'SUPPLIER CHARGES', 'TRANSPORT',
            'FOOD EXPENSES', 'OTHERS', 'TOTAL', 'PO Number'
        ]
        for col_idx, h in enumerate(headers, 1):
            ws.cell(2, col_idx, h)

        # Row 1
        ws.cell(3, 1, 1)
        ws.cell(3, 2, "08-10-2026")
        ws.cell(3, 3, "Subbaramireddy")
        ws.cell(3, 4, "N. NAGENDRA")
        ws.cell(3, 5, "Rahamtulla")
        ws.cell(3, 6, "Naidupeta")
        ws.cell(3, 7, "Pyraxalt")
        ws.cell(3, 8, "PADDY")
        ws.cell(3, 9, "LFM")
        ws.cell(3, 10, "Verubotlaplli")
        ws.cell(3, 11, 55)   # Farmers (should NOT be in expense total)
        ws.cell(3, 12, 325)  # Supplier Charges
        ws.cell(3, 13, 875)  # Transport
        ws.cell(3, 14, 825)  # Food Expenses
        ws.cell(3, 15, 2475) # Others
        ws.cell(3, 16, 4500) # Total (325 + 875 + 825 + 2475 = 4500)
        ws.cell(3, 17, "4800108505")

        acts = extract_activities_from_sheet(ws)
        self.assertEqual(len(acts), 1)
        act = acts[0]

        # Verify Supplier Charges was mapped to tent and farmers was preserved
        self.assertEqual(act['farmers'], 55)
        self.assertEqual(act['tent'], 325.0)
        self.assertEqual(act['transport'], 875.0)
        self.assertEqual(act['food'], 825.0)
        self.assertEqual(act['others'], 2475.0)
        # Sum of all activities expenses other than no of farmers:
        exp_sum = act['tent'] + act['transport'] + act['food'] + act['others']
        self.assertEqual(exp_sum, 4500.0)

    def test_format_tbm_workbook_step1(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            fpath = os.path.join(tmpdir, "test_tbm.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            ws.cell(1, 1, "Activities Expenses by Naidupeta Tbm Nagendra")
            headers = [
                'S.NO', 'DATE', 'ZDGM', 'TBM', 'MDO', 'TERRITORY', 'PRODUCT', 'CROP',
                'ACTIVITY', 'VILLAGE', 'NO.OF FARMERS', 'SUPPLIER CHARGES', 'TRANSPORT',
                'FOOD EXPENSES', 'OTHERS', 'TOTAL', 'PO Number'
            ]
            for col_idx, h in enumerate(headers, 1):
                ws.cell(2, col_idx, h)

            ws.cell(3, 1, 1)
            ws.cell(3, 2, "08-10-2026")
            ws.cell(3, 3, "Subbaramireddy")
            ws.cell(3, 4, "N. NAGENDRA")
            ws.cell(3, 5, "Rahamtulla")
            ws.cell(3, 6, "Naidupeta")
            ws.cell(3, 7, "Pyraxalt")
            ws.cell(3, 8, "PADDY")
            ws.cell(3, 9, "LFM")
            ws.cell(3, 10, "Verubotlaplli")
            ws.cell(3, 11, 55)
            ws.cell(3, 12, 325)
            ws.cell(3, 13, 875)
            ws.cell(3, 14, 825)
            ws.cell(3, 15, 2475)
            ws.cell(3, 16, 4500)
            ws.cell(3, 17, "4800108505")

            wb.save(fpath)

            res = format_tbm_workbook(fpath)
            self.assertTrue(res['success'])

            wb_res = openpyxl.load_workbook(fpath)
            self.assertIn("Sheet2", wb_res.sheetnames)
            ws2 = wb_res["Sheet2"]

            # Row 3 in Sheet 2:
            # Col 11: Farmers = 55
            # Col 12: Tent/Hall Suppliers Charges = 325
            # Col 13: Food Expenses = 825
            # Col 14: Transport = 875
            # Col 15: Others/Gifts = 2475
            # Col 16: Total = =SUM(L3:O3)
            self.assertEqual(ws2.cell(3, 11).value, 55)
            self.assertEqual(ws2.cell(3, 12).value, 325)
            self.assertEqual(ws2.cell(3, 13).value, 825)
            self.assertEqual(ws2.cell(3, 14).value, 875)
            self.assertEqual(ws2.cell(3, 15).value, 2475)
            self.assertEqual(ws2.cell(3, 16).value, "=SUM(L3:O3)")

    def test_formula_evaluation_and_expense_columns(self):
        """Tests that formulas like =K3*75 in food expenses and =SUM(L3:O3) in total are evaluated."""
        with tempfile.TemporaryDirectory() as tmpdir:
            fpath = os.path.join(tmpdir, "test_formula_tbm.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            ws.cell(1, 1, "Activities Expenses by Nellore - 2 TBM Manish")
            headers = [
                'S.No', 'Date', 'ZDGM', 'TBM', 'TBM Territory', 'MDO', 'Product', 'Crop',
                'Activity', 'Village', 'No.of Farmers', 'Tent/Hall Suppliers Charges',
                'Food Expenses', 'Transport', 'Others', 'Total', 'PO Number'
            ]
            for col_idx, h in enumerate(headers, 1):
                ws.cell(2, col_idx, h)

            # Row 1 (Row 3 in sheet)
            ws.cell(3, 1, 1)
            ws.cell(3, 2, "19/8/26")
            ws.cell(3, 3, "Subbarami Reddy")
            ws.cell(3, 4, "Manish")
            ws.cell(3, 5, "Nellore - 2")
            ws.cell(3, 6, "Salman")
            ws.cell(3, 7, "Pyraxalt")
            ws.cell(3, 8, "Paddy")
            ws.cell(3, 9, "MFM")
            ws.cell(3, 10, "Pallipalem")
            ws.cell(3, 11, 80)          # Col K: Farmers
            ws.cell(3, 12, 2500)        # Col L: Tent
            ws.cell(3, 13, "=K3*75")    # Col M: Food = 80 * 75 = 6000
            ws.cell(3, 14, 1000)        # Col N: Transport
            ws.cell(3, 15, None)        # Col O: Others
            ws.cell(3, 16, "=SUM(L3:O3)") # Col P: Total = 9500
            ws.cell(3, 17, "4800108505")

            # Row 2 (Row 4 in sheet)
            ws.cell(4, 1, 2)
            ws.cell(4, 2, "22/8/26")
            ws.cell(4, 3, "Subbarami Reddy")
            ws.cell(4, 4, "Manish")
            ws.cell(4, 5, "Nellore - 2")
            ws.cell(4, 6, "Tirumala Naidu")
            ws.cell(4, 7, "Pyraxalt")
            ws.cell(4, 8, "Paddy")
            ws.cell(4, 9, "MFM")
            ws.cell(4, 10, "Uppalapadu")
            ws.cell(4, 11, 75)          # Col K: Farmers
            ws.cell(4, 12, 800)         # Col L: Tent
            ws.cell(4, 13, "=K4*60")    # Col M: Food = 75 * 60 = 4500
            ws.cell(4, 14, 600)         # Col N: Transport
            ws.cell(4, 15, None)        # Col O: Others
            ws.cell(4, 16, "=SUM(L4:O4)") # Col P: Total = 5900
            ws.cell(4, 17, "4800108505")

            wb.save(fpath)

            acts = extract_activities_from_sheet(ws)
            self.assertEqual(len(acts), 2)
            self.assertEqual(acts[0]['farmers'], 80)
            self.assertEqual(acts[0]['tent'], 2500.0)
            self.assertEqual(acts[0]['food'], 6000.0)
            self.assertEqual(acts[0]['transport'], 1000.0)
            self.assertEqual(acts[0]['others'], 0.0)
            self.assertEqual(acts[0]['total'], 9500.0)

            self.assertEqual(acts[1]['farmers'], 75)
            self.assertEqual(acts[1]['tent'], 800.0)
            self.assertEqual(acts[1]['food'], 4500.0)
            self.assertEqual(acts[1]['transport'], 600.0)
            self.assertEqual(acts[1]['total'], 5900.0)

            res = format_tbm_workbook(fpath)
            self.assertTrue(res['success'])

            wb_out = openpyxl.load_workbook(fpath, data_only=False)
            ws2 = wb_out['Sheet2']
            # First row of data in Sheet2 is row 3
            self.assertEqual(ws2.cell(3, 11).value, 80)
            self.assertEqual(ws2.cell(3, 12).value, 2500.0)
            self.assertEqual(ws2.cell(3, 13).value, 6000.0)
            self.assertEqual(ws2.cell(3, 14).value, 1000.0)
            self.assertEqual(ws2.cell(3, 16).value, "=SUM(L3:O3)")

            # Second row of data in Sheet2 is row 4
            self.assertEqual(ws2.cell(4, 11).value, 75)
            self.assertEqual(ws2.cell(4, 12).value, 800.0)
            self.assertEqual(ws2.cell(4, 13).value, 4500.0)
            self.assertEqual(ws2.cell(4, 14).value, 600.0)
            self.assertEqual(ws2.cell(4, 16).value, "=SUM(L4:O4)")

    def test_marking_as_done_and_skipping(self):
        """Tests that formatting marks Sheet 1 Column R as DONE and skips already formatted files."""
        with tempfile.TemporaryDirectory() as tmpdir:
            fpath = os.path.join(tmpdir, "test_skip_tbm.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            ws.cell(1, 1, "Activities Expenses by Kandukur Tbm Govind")
            headers = [
                'Sl.No', 'Date', 'Zdgm', 'TBM', 'MDO', 'Territory', 'Product', 'Crop',
                'Activity', 'Village', 'No.of Farmers', 'Tent/Hall Suppliers charges',
                'Food Expenses', 'Transport', 'Others', 'Total', 'PO Number'
            ]
            for col_idx, h in enumerate(headers, 1):
                ws.cell(2, col_idx, h)

            ws.cell(3, 1, 1)
            ws.cell(3, 2, "13/7/26")
            ws.cell(3, 3, "Subbaramireddy")
            ws.cell(3, 4, "Govind")
            ws.cell(3, 5, "V Y REDDY")
            ws.cell(3, 6, "KANDUKURU")
            ws.cell(3, 7, "Pyraxalt")
            ws.cell(3, 8, "PADDY")
            ws.cell(3, 9, "LFM")
            ws.cell(3, 10, "guduluru")
            ws.cell(3, 11, 50)
            ws.cell(3, 12, 1000)
            ws.cell(3, 13, 2700)
            ws.cell(3, 14, 800)
            ws.cell(3, 16, 4500)
            ws.cell(3, 17, "4800108505")

            wb.save(fpath)

            # First run: should format and mark as DONE
            res1 = format_tbm_workbook(fpath)
            self.assertTrue(res1['success'])
            self.assertFalse(res1.get('skipped', False))

            # Inspect Sheet 1: Column R (Col 18) must have 'Status' in row 2 and 'DONE' in row 3
            wb_check = openpyxl.load_workbook(fpath)
            ws1 = wb_check["Sheet1"]
            self.assertEqual(ws1.cell(2, 18).value, "Status")
            self.assertEqual(ws1.cell(3, 18).value, "DONE")
            wb_check.close()

            # Second run (force=False): should skip because it is already marked as DONE
            res2 = format_tbm_workbook(fpath, force=False)
            self.assertTrue(res2['success'])
            self.assertTrue(res2.get('skipped', False))
            self.assertTrue(res2.get('alreadyFormatted', False))

            # Third run (force=True): should re-format even though already done
            res3 = format_tbm_workbook(fpath, force=True)
            self.assertTrue(res3['success'])
            self.assertFalse(res3.get('skipped', False))

    def test_unformat_when_done_removed(self):
        """
        Tests that when a user manually deletes/removes 'DONE' from Sheet 1,
        the system treats the sheet as NOT formatted (even though Sheet2 exists)
        and re-processes it purely based on the presence of DONE in Sheet 1.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            fpath = os.path.join(tmpdir, "test_unformat.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            ws.cell(1, 1, "Activities Expenses by Kandukur Tbm Govind")
            headers = [
                'Sl.No', 'Date', 'Zdgm', 'TBM', 'MDO', 'Territory', 'Product', 'Crop',
                'Activity', 'Village', 'No.of Farmers', 'Tent/Hall Suppliers charges',
                'Food Expenses', 'Transport', 'Others', 'Total', 'PO Number'
            ]
            for col_idx, h in enumerate(headers, 1):
                ws.cell(2, col_idx, h)

            ws.cell(3, 1, 1)
            ws.cell(3, 2, "13/7/26")
            ws.cell(3, 3, "Subbaramireddy")
            ws.cell(3, 4, "Govind")
            ws.cell(3, 5, "V Y REDDY")
            ws.cell(3, 6, "KANDUKURU")
            ws.cell(3, 7, "Pyraxalt")
            ws.cell(3, 8, "PADDY")
            ws.cell(3, 9, "LFM")
            ws.cell(3, 10, "guduluru")
            ws.cell(3, 11, 50)
            ws.cell(3, 12, 1000)
            ws.cell(3, 13, 2700)
            ws.cell(3, 14, 800)
            ws.cell(3, 16, 4500)
            ws.cell(3, 17, "4800108505")

            wb.save(fpath)

            # 1. Format first time
            res1 = format_tbm_workbook(fpath)
            self.assertTrue(res1['success'])
            self.assertFalse(res1['skipped'])

            # 2. Simulate user opening Excel, removing 'DONE' from Column R, but leaving Sheet2
            wb_mod = openpyxl.load_workbook(fpath)
            ws1_mod = wb_mod["Sheet1"]
            self.assertEqual(ws1_mod.cell(3, 18).value, "DONE")
            # Clear all DONE cells in Sheet 1
            for r in range(1, ws1_mod.max_row + 1):
                for c in range(1, ws1_mod.max_column + 1):
                    if str(ws1_mod.cell(r, c).value or "").strip().upper() == "DONE":
                        ws1_mod.cell(r, c).value = None
            wb_mod.save(fpath)
            wb_mod.close()

            # 3. System re-runs with force=False: MUST NOT SKIP, MUST PROCESS PURELY ON LOGIC
            res2 = format_tbm_workbook(fpath, force=False)
            self.assertTrue(res2['success'])
            self.assertFalse(res2.get('skipped', False))
            self.assertFalse(res2.get('alreadyFormatted', False))

            # 4. Sheet 1 must now have DONE restored
            wb_check = openpyxl.load_workbook(fpath)
            self.assertEqual(wb_check["Sheet1"].cell(3, 18).value, "DONE")
            wb_check.close()

    def test_step2_consolidation_and_sheet2_done_marking(self):
        """
        Tests that Step 2:
        1. Consolidates formatted tables from Sheet 2 into Master Summary.
        2. Marks Sheet 2 in the source file as DONE (Column 18 / Status column).
        3. Skips already marked Sheet 2 tables on subsequent runs.
        4. Re-consolidates when DONE is removed from Sheet 2 (pure logic).
        5. Re-consolidates when force=True.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            tbm_folder = os.path.join(tmpdir, "Govind")
            os.makedirs(tbm_folder, exist_ok=True)
            fpath = os.path.join(tbm_folder, "Govind_bills.xlsx")
            master_output = os.path.join(tmpdir, "Test-All-TBMs-Summary.xlsx")

            wb = openpyxl.Workbook()
            ws1 = wb.active
            ws1.title = "Sheet1"

            ws1.cell(1, 1, "Activities Expenses by Kandukur Tbm Govind")
            headers = [
                'Sl.No', 'Date', 'Zdgm', 'TBM', 'MDO', 'Territory', 'Product', 'Crop',
                'Activity', 'Village', 'No.of Farmers', 'Tent/Hall Suppliers charges',
                'Food Expenses', 'Transport', 'Others', 'Total', 'PO Number'
            ]
            for col_idx, h in enumerate(headers, 1):
                ws1.cell(2, col_idx, h)

            ws1.cell(3, 1, 1)
            ws1.cell(3, 2, "13/7/26")
            ws1.cell(3, 3, "Subbaramireddy")
            ws1.cell(3, 4, "Govind")
            ws1.cell(3, 5, "V Y REDDY")
            ws1.cell(3, 6, "KANDUKURU")
            ws1.cell(3, 7, "Pyraxalt")
            ws1.cell(3, 8, "PADDY")
            ws1.cell(3, 9, "LFM")
            ws1.cell(3, 10, "guduluru")
            ws1.cell(3, 11, 50)
            ws1.cell(3, 12, 1000)
            ws1.cell(3, 13, 2700)
            ws1.cell(3, 14, 800)
            ws1.cell(3, 16, 4500)
            ws1.cell(3, 17, "4800108505")

            wb.save(fpath)

            # Step 1: Format workbook (generates Sheet 2)
            res1 = format_tbm_workbook(fpath)
            self.assertTrue(res1['success'])

            # Verify Sheet 2 exists and Col 18 does NOT have DONE yet
            wb_src = openpyxl.load_workbook(fpath)
            self.assertIn("Sheet2", wb_src.sheetnames)
            ws2 = wb_src["Sheet2"]
            self.assertNotEqual(ws2.cell(3, 18).value, "DONE")
            wb_src.close()

            # Step 2: Consolidate Master Summary (1st run)
            res2 = generate_tbm_summary(tmpdir, output_path=master_output, force=False)
            self.assertTrue(res2['success'])
            self.assertEqual(res2['totalActivities'], 1)
            self.assertTrue(os.path.exists(master_output))

            # Verify Sheet 2 now HAS 'Status' in row 2 and 'DONE' in row 3 (Col 18)
            wb_src2 = openpyxl.load_workbook(fpath)
            ws2_check = wb_src2["Sheet2"]
            self.assertEqual(ws2_check.cell(2, 18).value, "Status")
            self.assertEqual(ws2_check.cell(3, 18).value, "DONE")
            wb_src2.close()

            # Step 2 (2nd run, force=False): Must skip already DONE tables!
            res3 = generate_tbm_summary(tmpdir, output_path=master_output, force=False)
            self.assertTrue(res3['success'])
            self.assertTrue(res3.get('alreadyConsolidated', False))
            self.assertEqual(res3['totalActivities'], 0)

            # Step 2 (3rd run, user manually removes DONE from Sheet 2): Pure logic re-consolidates!
            wb_mod = openpyxl.load_workbook(fpath)
            ws2_mod = wb_mod["Sheet2"]
            ws2_mod.cell(3, 18).value = None
            wb_mod.save(fpath)
            wb_mod.close()

            res4 = generate_tbm_summary(tmpdir, output_path=master_output, force=False)
            self.assertTrue(res4['success'])
            self.assertFalse(res4.get('alreadyConsolidated', False))
            self.assertEqual(res4['totalActivities'], 1)

            # Verify Sheet 2 has DONE restored
            wb_src3 = openpyxl.load_workbook(fpath)
            self.assertEqual(wb_src3["Sheet2"].cell(3, 18).value, "DONE")
            wb_src3.close()

            # Step 2 (4th run, force=True): Consolidates even when DONE is present
            res5 = generate_tbm_summary(tmpdir, output_path=master_output, force=True)
            self.assertTrue(res5['success'])
            self.assertEqual(res5['totalActivities'], 1)


if __name__ == '__main__':
    unittest.main()
