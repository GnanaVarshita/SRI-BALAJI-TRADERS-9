"""
Unit Tests for Card Sync Engine (backend.card_sync)
Sri Balaji Traders Automation System
"""

import unittest
import tempfile
import shutil
from pathlib import Path
import openpyxl

from backend.card_sync.constants import (
    normalize_activity,
    match_activity_to_column,
)
from backend.card_sync.extractor import (
    find_summary_files,
    extract_all_tbms_spent,
    mark_tbms_as_synced,
)
from backend.card_sync.corteva_sync import (
    sync_corteva_cards_workbook,
    is_corteva_row_occupied,
)
from backend.card_sync.fmc_sync import (
    sync_fmc_cards_workbook,
    is_fmc_row_occupied,
)
from backend.card_sync.engine import sync_tbm_with_cards


class TestCardSyncEngine(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _create_sample_master_summary(self, file_path, po="4800108505", tbm="Govind", rows_count=3):
        """Creates a sample All-TBMs summary Excel matching screenshot ss 1."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Table title at row 5
        ws.cell(5, 1, f"MARKETING ACTIVITIES EXPENSES-{tbm.upper()}-KANDUKURU")

        # Headers at row 6
        headers = [
            "SI No", "Date", "ZDGM", "TBM", "MDO", "Territory",
            "Product", "Crop", "Activity", "Village", "No. of Farmers",
            "Tent/Hall/Chairs Suppliers", "Food Expenses", "Transport", "Others/Gifts",
            "Total", "PO Number"
        ]
        for c_idx, h in enumerate(headers, 1):
            ws.cell(6, c_idx, h)

        # Data rows
        for i in range(rows_count):
            r = 7 + i
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, "15-07-2026")
            ws.cell(r, 3, "Subbaramireddy")
            ws.cell(r, 4, tbm)
            ws.cell(r, 5, "V Y REDDY")
            ws.cell(r, 6, "KANDUKURU")
            ws.cell(r, 7, "Pyraxalt")
            ws.cell(r, 8, "PADDY")
            ws.cell(r, 9, "field day")
            ws.cell(r, 10, "Machavaram")
            ws.cell(r, 11, 30)
            ws.cell(r, 13, 1000)
            ws.cell(r, 16, 1000)  # Total
            ws.cell(r, 17, po)     # PO Number

        tot_r = 7 + rows_count
        ws.cell(tot_r, 1, "TOTAL")
        ws.cell(tot_r, 16, rows_count * 1000)

        wb.save(file_path)
        wb.close()

    def _create_sample_corteva_card(self, file_path, po="4800108505"):
        """Creates a sample Corteva card workbook with an existing data row at row 12."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pyraxalt (MA)"

        # PO in cell(6, 1)
        ws.cell(6, 1, po)
        ws.cell(7, 4, "Nellore - Kandukuru")

        # Activity headers in row 10
        ws.cell(10, 12, "MFM(RGL)")
        ws.cell(10, 13, "LFM(RGL)")
        ws.cell(10, 14, "FD")
        ws.cell(10, 15, "Amount")

        # Rate grid at T3:U5
        ws.cell(3, 20, "MFM(RGL)")
        ws.cell(3, 21, 100000)
        ws.cell(4, 20, "LFM(RGL)")
        ws.cell(4, 21, 50000)
        ws.cell(5, 20, "FD")
        ws.cell(5, 21, 30000)

        # Existing row from previous sync at row 12
        ws.cell(12, 1, 45)  # I.V NO
        ws.cell(12, 2, "2026-07-01")
        ws.cell(12, 3, "Nellore")
        ws.cell(12, 4, po)
        ws.cell(12, 5, "Marketing")
        ws.cell(12, 6, "Pyraxalt")
        ws.cell(12, 7, "Paddy")
        ws.cell(12, 8, "MFM")
        ws.cell(12, 10, "Prior TBM")
        ws.cell(12, 11, 2)
        ws.cell(12, 12, 5000)  # MFM amount
        ws.cell(12, 15, "=SUM(L12:N12)")

        wb.save(file_path)
        wb.close()

    def test_activity_normalization(self):
        """Test normalization and matching of activity names."""
        self.assertEqual(normalize_activity("FIELD DAYS"), "FD")
        self.assertEqual(normalize_activity("field day"), "FD")
        self.assertEqual(normalize_activity("LFM(RGL)"), "LFM")
        self.assertEqual(normalize_activity("Large Farmer Meeting (RGL)"), "LFM")
        self.assertEqual(normalize_activity("MEGA FARMER MEETING / NEW PRODUCT LAUCHES(RGL)"), "MFM")
        self.assertEqual(normalize_activity("G.M (RGL)"), "GM_RGL")
        self.assertEqual(normalize_activity("G.M (Regular)"), "GM_REGULAR")

        col_map = {"MFM": 12, "LFM": 13, "FD": 14}
        self.assertEqual(match_activity_to_column("field day", col_map), 14)
        self.assertEqual(match_activity_to_column("LFM(RGL)", col_map), 13)
        self.assertEqual(match_activity_to_column("MEGA FARMER MEETING", col_map), 12)

    def test_extract_and_mark_all_tbms_spent(self):
        """Test extraction from All-TBMs summary file, marking DONE, and skipping on subsequent run."""
        summary_path = self.temp_dir / "All-TBMs-Summary.xlsx"
        self._create_sample_master_summary(summary_path, po="4800108505", tbm="Govind", rows_count=3)

        # 1. Initial extraction
        tbm_data, items_to_mark, skipped_files, total_files = extract_all_tbms_spent(summary_path)
        self.assertEqual(total_files, 1)
        self.assertEqual(skipped_files, 0)
        self.assertIn("4800108505", tbm_data)
        self.assertEqual(len(tbm_data["4800108505"]), 1)
        item = tbm_data["4800108505"][0]
        self.assertEqual(item["tbm"], "Govind")
        self.assertEqual(item["num_activities"], 3)
        self.assertEqual(item["total_amount"], 3000)

        # 2. Mark as DONE
        marked_cnt, mark_errors = mark_tbms_as_synced(items_to_mark)
        self.assertEqual(marked_cnt, 1)
        self.assertEqual(len(mark_errors), 0)

        # Verify Column R in the saved file
        wb_check = openpyxl.load_workbook(summary_path)
        ws_check = wb_check["Sheet1"]
        self.assertEqual(ws_check.cell(6, 18).value, "Status")
        self.assertEqual(ws_check.cell(7, 18).value, "DONE")
        self.assertEqual(ws_check.cell(8, 18).value, "DONE")
        self.assertEqual(ws_check.cell(9, 18).value, "DONE")
        wb_check.close()

        # 3. Second extraction without force: should be skipped
        tbm_data_2, items_to_mark_2, skipped_files_2, total_files_2 = extract_all_tbms_spent(summary_path)
        self.assertEqual(len(tbm_data_2), 0)
        self.assertEqual(skipped_files_2, 1)

        # 4. User clears DONE from row 7 ("pure logic"): row 7 should be extracted again
        wb_edit = openpyxl.load_workbook(summary_path)
        wb_edit["Sheet1"].cell(7, 18).value = None
        wb_edit.save(summary_path)
        wb_edit.close()

        tbm_data_3, _, skipped_files_3, _ = extract_all_tbms_spent(summary_path)
        self.assertEqual(skipped_files_3, 0)
        self.assertIn("4800108505", tbm_data_3)
        self.assertEqual(tbm_data_3["4800108505"][0]["num_activities"], 1)

    def test_non_destructive_corteva_sync(self):
        """Test that existing rows in Corteva cards are preserved and new rows are appended."""
        card_path = self.temp_dir / "June 2 Nellore.xlsx"
        self._create_sample_corteva_card(card_path, po="4800108505")

        # Fake TBM data to sync
        tbm_data = {
            "4800108505": [
                {
                    "po": "4800108505",
                    "tbm": "Govind",
                    "product": "Pyraxalt",
                    "crop": "Paddy",
                    "activity": "field day",
                    "num_activities": 3,
                    "total_amount": 3000,
                    "area": "Kandukuru",
                    "zdgm": "Subbaramireddy"
                }
            ]
        }

        wb = openpyxl.load_workbook(card_path)
        updated_count = sync_corteva_cards_workbook(wb, tbm_data, service_charge_percent=5.0)
        self.assertEqual(updated_count, 1)

        ws = wb["Pyraxalt (MA)"]
        # Verify row 12 (prior sync) was NOT overwritten or cleared
        self.assertEqual(ws.cell(12, 1).value, 45)
        self.assertEqual(ws.cell(12, 10).value, "Prior TBM")
        self.assertEqual(ws.cell(12, 12).value, 5000)

        # Verify row 13 has the newly appended spending
        self.assertEqual(ws.cell(13, 10).value, "Govind")
        self.assertEqual(ws.cell(13, 8).value, "field day")
        self.assertEqual(ws.cell(13, 14).value, 3000)  # Col 14 is FD
        self.assertEqual(ws.cell(13, 15).value, "=SUM(L13,M13,N13)")  # Row total in Amount col

        # Verify row 29 sums both rows
        self.assertEqual(ws.cell(29, 12).value, "=SUM(L12:L28)")
        self.assertEqual(ws.cell(29, 14).value, "=SUM(N12:N28)")

        # Verify Right Summary table formulas (Col V = SPENT, Col W = SV Charges, Col X = Total IV, Col Y = Balance)
        # FD is row 18 (16 + 2)
        self.assertEqual(ws.cell(18, 22).value, "=N29")
        self.assertEqual(ws.cell(18, 23).value, "=V18*0.05")
        self.assertEqual(ws.cell(18, 24).value, "=V18+W18")
        self.assertEqual(ws.cell(18, 25).value, "=U18-X18")

        wb.close()

    def test_sync_with_folder_of_summaries(self):
        """Test full sync_tbm_with_cards when passed a folder containing multiple summary Excels."""
        summaries_folder = self.temp_dir / "All-TBMs-Summaries"
        summaries_folder.mkdir()

        # Summary 1 has Govind
        self._create_sample_master_summary(
            summaries_folder / "Summary1.xlsx", po="4800108505", tbm="Govind", rows_count=2
        )
        # Summary 2 has Manish
        self._create_sample_master_summary(
            summaries_folder / "Summary2.xlsx", po="4800108505", tbm="Manish", rows_count=3
        )

        card_path = self.temp_dir / "June 2 Nellore.xlsx"
        self._create_sample_corteva_card(card_path, po="4800108505")

        # Run sync passing the folder path
        res = sync_tbm_with_cards(
            cards_excel_path=card_path,
            tbm_summary_excel_path=summaries_folder,
            service_charge_percent=5.0
        )

        self.assertTrue(res["success"])
        self.assertEqual(res["updatedCards"], 1)
        self.assertEqual(res["totalTbmPOs"], 1)
        self.assertEqual(res["syncedRecords"], 5)  # 2 from Summary1 + 3 from Summary2
        self.assertEqual(res["markedFilesCount"], 2)

        # Verify both summary files were marked as DONE
        for fname in ["Summary1.xlsx", "Summary2.xlsx"]:
            wb_s = openpyxl.load_workbook(summaries_folder / fname)
            ws_s = wb_s["Sheet1"]
            self.assertEqual(ws_s.cell(6, 18).value, "Status")
            self.assertEqual(ws_s.cell(7, 18).value, "DONE")
            wb_s.close()

        # Verify PO Card has row 12 (prior), row 13 (Govind), row 14 (Manish)
        wb_c = openpyxl.load_workbook(card_path)
        ws_c = wb_c["Pyraxalt (MA)"]
        self.assertEqual(ws_c.cell(12, 10).value, "Prior TBM")
        self.assertEqual(ws_c.cell(13, 10).value, "Govind")
        self.assertEqual(ws_c.cell(14, 10).value, "Manish")
        wb_c.close()

        # Re-running sync on the same folder should skip all already-done files
        res2 = sync_tbm_with_cards(
            cards_excel_path=card_path,
            tbm_summary_excel_path=summaries_folder,
            service_charge_percent=5.0
        )
        self.assertTrue(res2["success"])
        self.assertEqual(res2["skippedFilesCount"], 2)
        self.assertEqual(res2["updatedCards"], 0)

    def test_corteva_right_summary_preservation_and_totals(self):
        """Test that Right Summary table (Cols T-Y) and Row 29 column sums are fully computed."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Herbi (MKTG)"

        # PO header
        ws.cell(6, 1, "4800110294")
        ws.cell(7, 4, "Bhaskar - Nellore")

        # Activity headers in row 10
        headers_10 = ["MFM (RGL)", 0, "FD", 0, 0, "BVC", None, "Amount"]
        for idx, h in enumerate(headers_10):
            ws.cell(10, 12 + idx, h)

        # Right Summary table at rows 15-22
        ws.cell(15, 20, "Activities")
        ws.cell(15, 21, "BUDGET")
        ws.cell(15, 22, "SPENT")
        ws.cell(15, 23, "SV Charges")
        ws.cell(15, 24, "TOTAL IV")
        ws.cell(15, 25, "BALANCE")

        activities = [
            ("MFM (RGL)", 180000),
            (0, 0),
            ("FD", 36000),
            (0, 0),
            (0, 0),
            ("BVC", 115500),
        ]
        for i, (act, bgt) in enumerate(activities):
            r = 16 + i
            ws.cell(r, 20, act)
            ws.cell(r, 21, bgt)

        ws.cell(22, 20, "TOTAL")

        # Existing row 12
        ws.cell(12, 1, 45)
        ws.cell(12, 12, 50000)

        tbm_data = {
            "4800110294": [
                {
                    "tbm": "Pradeep",
                    "product": "Herbicides",
                    "crop": "Paddy",
                    "activity": "FD",
                    "num_activities": 2,
                    "total_amount": 10000,
                    "area": "Nellore",
                    "zdgm": "Bhaskar"
                }
            ]
        }

        updated = sync_corteva_cards_workbook(wb, tbm_data, service_charge_percent=5.0)
        self.assertEqual(updated, 1)

        # Row 29 column sums
        self.assertEqual(ws.cell(29, 12).value, "=SUM(L12:L28)")
        self.assertEqual(ws.cell(29, 14).value, "=SUM(N12:N28)")
        self.assertEqual(ws.cell(29, 17).value, "=SUM(Q12:Q28)")
        self.assertEqual(ws.cell(29, 19).value, "=SUM(S12:S28)")

        # Right Summary rows 16..21
        self.assertEqual(ws.cell(16, 22).value, "=L29")
        self.assertEqual(ws.cell(16, 23).value, "=V16*0.05")
        self.assertEqual(ws.cell(16, 24).value, "=V16+W16")
        self.assertEqual(ws.cell(16, 25).value, "=U16-X16")

        self.assertEqual(ws.cell(18, 22).value, "=N29")
        self.assertEqual(ws.cell(18, 23).value, "=V18*0.05")

        # BVC in row 21: SV Charges should remain None when not present initially
        self.assertEqual(ws.cell(21, 22).value, "=Q29")
        self.assertIsNone(ws.cell(21, 23).value)
        self.assertEqual(ws.cell(21, 24).value, "=V21+W21")
        self.assertEqual(ws.cell(21, 25).value, "=U21-X21")

        # TOTAL in row 22
        self.assertEqual(ws.cell(22, 20).value, "TOTAL")
        self.assertEqual(ws.cell(22, 21).value, "=SUM(U16:U21)")
        self.assertEqual(ws.cell(22, 22).value, "=SUM(V16:V21)")
        self.assertEqual(ws.cell(22, 23).value, "=SUM(W16:W21)")
        self.assertEqual(ws.cell(22, 24).value, "=SUM(X16:X21)")
        self.assertEqual(ws.cell(22, 25).value, "=SUM(Y16:Y21)")

        wb.close()


if __name__ == "__main__":
    unittest.main()
