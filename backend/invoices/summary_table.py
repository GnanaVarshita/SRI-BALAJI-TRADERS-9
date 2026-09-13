"""
Invoice Details Cumulative Summary Table Renderer
Sri Balaji Traders Automation System
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def render_details_cumulative_summary_table(ws, start_row, summary_items, service_charge_pct=5.0, styles=None):
    """
    Renders a summary table at the end of Sheet2 displaying cumulative values of all products & activities,
    vendor charges (e.g. 5%), and total values, followed by a GRAND TOTAL row.
    
    Columns used: F (S.No), G (Product), H (Crop), I (Activity), J (Cumulative Value),
                  K (Vendor Charges), L (Total Amount).
    """
    if not summary_items:
        return start_row

    # Styles
    thin_black = Side(style='thin', color='000000')
    double_black = Side(style='double', color='000000')
    box_border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    grand_border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=double_black)

    title_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    grand_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')

    font_title = Font(name='Calibri', size=11, bold=True, color='006100')
    font_header = Font(name='Calibri', size=10, bold=True, color='006100')
    font_bold = Font(name='Calibri', size=10, bold=True)
    font_regular = Font(name='Calibri', size=10)
    font_grand = Font(name='Calibri', size=10, bold=True, color='9C0006')

    align_center = Alignment(horizontal='center', vertical='center')
    align_center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    # Ensure column widths
    ws.column_dimensions['K'].width = max(ws.column_dimensions['K'].width or 12, 18)
    ws.column_dimensions['L'].width = max(ws.column_dimensions['L'].width or 15, 18)

    current_r = start_row

    # 1. Title Banner across Columns F to L (cols 6 to 12)
    ws.row_dimensions[current_r].height = 24
    ws.merge_cells(start_row=current_r, start_column=6, end_row=current_r, end_column=12)
    title_c = ws.cell(current_r, 6, "CUMULATIVE EXPENSES & VENDOR CHARGES SUMMARY")
    title_c.font = font_title
    title_c.alignment = align_center
    title_c.fill = title_fill

    for c in range(6, 13):
        ws.cell(current_r, c).border = box_border

    current_r += 1

    # 2. Table Column Headers
    ws.row_dimensions[current_r].height = 26
    headers = [
        (6, "S.No", align_center),
        (7, "Product", align_center),
        (8, "Crop", align_center),
        (9, "Activity", align_center),
        (10, "Cumulative Value", align_right),
        (11, f"Vendor Charges ({service_charge_pct:g}%)", align_right),
        (12, "Total Amount", align_right)
    ]

    for col_idx, h_text, al in headers:
        cell = ws.cell(current_r, col_idx, h_text)
        cell.font = font_header
        cell.alignment = align_center_wrap if col_idx in [10, 11, 12] else al
        cell.fill = header_fill
        cell.border = box_border

    current_r += 1
    start_data_r = current_r

    # 3. Data Rows
    for idx, item in enumerate(summary_items, start=1):
        ws.row_dimensions[current_r].height = 20

        # S.No
        c_sno = ws.cell(current_r, 6, idx)
        c_sno.font = font_bold
        c_sno.alignment = align_center
        c_sno.border = box_border

        # Product
        c_prod = ws.cell(current_r, 7, item.get('product', ''))
        c_prod.font = font_regular
        c_prod.alignment = align_center
        c_prod.border = box_border

        # Crop
        c_crop = ws.cell(current_r, 8, item.get('crop', ''))
        c_crop.font = font_regular
        c_crop.alignment = align_center
        c_crop.border = box_border

        # Activity
        c_act = ws.cell(current_r, 9, item.get('activity', ''))
        c_act.font = font_bold
        c_act.alignment = align_center
        c_act.border = box_border

        # Cumulative Value (Formula link if present, else number)
        sub_ref = item.get('subtot_ref')
        val_j = f"={sub_ref}" if sub_ref else item.get('raw_amount', 0.0)
        c_j = ws.cell(current_r, 10, val_j)
        c_j.font = font_bold
        c_j.alignment = align_right
        c_j.number_format = '#,##0.00'
        c_j.border = box_border

        # Vendor Charges
        sc_ref = item.get('sc_ref')
        val_k = f"={sc_ref}" if sc_ref else item.get('vendor_charges', 0.0)
        c_k = ws.cell(current_r, 11, val_k)
        c_k.font = font_bold
        c_k.alignment = align_right
        c_k.number_format = '#,##0.00'
        c_k.border = box_border

        # Total Amount
        tot_ref = item.get('tot_sc_ref')
        val_l = f"={tot_ref}" if tot_ref else item.get('total_amount', 0.0)
        c_l = ws.cell(current_r, 12, val_l)
        c_l.font = font_bold
        c_l.alignment = align_right
        c_l.number_format = '#,##0.00'
        c_l.border = box_border

        current_r += 1

    end_data_r = current_r - 1

    # 4. GRAND TOTAL Row
    r_grand = current_r
    ws.row_dimensions[r_grand].height = 24
    ws.merge_cells(start_row=r_grand, start_column=6, end_row=r_grand, end_column=9)
    lbl_g = ws.cell(r_grand, 6, "GRAND TOTAL")
    lbl_g.font = font_grand
    lbl_g.alignment = align_center
    lbl_g.fill = grand_fill

    # Sum of Cumulative Values
    c_tot_j = ws.cell(r_grand, 10, f"=SUM(J{start_data_r}:J{end_data_r})")
    c_tot_j.font = font_grand
    c_tot_j.alignment = align_right
    c_tot_j.number_format = '#,##0.00'
    c_tot_j.fill = grand_fill

    # Sum of Vendor Charges
    c_tot_k = ws.cell(r_grand, 11, f"=SUM(K{start_data_r}:K{end_data_r})")
    c_tot_k.font = font_grand
    c_tot_k.alignment = align_right
    c_tot_k.number_format = '#,##0.00'
    c_tot_k.fill = grand_fill

    # Sum of Total Amounts
    c_tot_l = ws.cell(r_grand, 12, f"=SUM(L{start_data_r}:L{end_data_r})")
    c_tot_l.font = font_grand
    c_tot_l.alignment = align_right
    c_tot_l.number_format = '#,##0.00'
    c_tot_l.fill = grand_fill

    for c in range(6, 13):
        ws.cell(r_grand, c).border = grand_border

    return r_grand + 2
