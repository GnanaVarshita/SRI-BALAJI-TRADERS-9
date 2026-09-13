"""
Corteva Invoice Sheet Builders
Sri Balaji Traders Automation System
"""

from openpyxl.styles import Font, Border, Side
from openpyxl.worksheet.pagebreak import Break
from common.formatters import clean_str
from common.currency import num_to_indian_words
from common.styles import setup_page_print_fit, write_sbt_header_block, write_bank_and_signature_block
try:
    from .summary_table import render_details_cumulative_summary_table
except ImportError:
    from invoices.summary_table import render_details_cumulative_summary_table


def render_corteva_invoice_block(ws, start_row, copy_type, invoice_no, invoice_date, po_number, metadata, activity_groups, service_charge_pct, styles):
    """
    Renders one full Corteva Tax Invoice block (ORIGINAL or DUPLICATE) starting at start_row.
    Returns: (end_row, r_subtotal, r_grand)
    """
    write_sbt_header_block(ws, start_row, copy_type, styles)

    # Customer & Metadata block (Rows start_row+8 to start_row+16)
    r_cust_start = start_row + 8
    for r_c in range(r_cust_start, r_cust_start + 9):
        ws.row_dimensions[r_c].height = 14

    ws[f'A{r_cust_start}'] = "TO:"
    ws[f'A{r_cust_start}'].font = styles['font_bold']
    ws[f'A{r_cust_start + 1}'] = "Corteva Agriscience India Private Limited"
    ws[f'A{r_cust_start + 1}'].font = styles['font_bold']
    ws[f'A{r_cust_start + 2}'] = "D.NO 12-112 Godown NO.2 Krishna Nagar"
    ws[f'A{r_cust_start + 2}'].font = styles['font_regular']
    ws[f'A{r_cust_start + 3}'] = "Vijayawada"
    ws[f'A{r_cust_start + 3}'].font = styles['font_regular']
    ws[f'A{r_cust_start + 4}'] = "01 520007"
    ws[f'A{r_cust_start + 4}'].font = styles['font_regular']
    ws[f'A{r_cust_start + 5}'] = "INDIA"
    ws[f'A{r_cust_start + 5}'].font = styles['font_regular']
    ws[f'A{r_cust_start + 6}'] = "GSTIN: 37AAACE2462M1ZI"
    ws[f'A{r_cust_start + 6}'].font = styles['font_bold']

    # Right metadata
    ws.merge_cells(f'I{r_cust_start}:J{r_cust_start}')
    ws[f'I{r_cust_start}'] = "MSME"
    ws[f'I{r_cust_start}'].font = styles['font_bold']
    ws[f'I{r_cust_start}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 1}'] = "Invoive no:"
    ws[f'G{r_cust_start + 1}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 1}:J{r_cust_start + 1}')
    ws[f'I{r_cust_start + 1}'] = invoice_no
    ws[f'I{r_cust_start + 1}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 1}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 2}'] = "Invoice Date:"
    ws[f'G{r_cust_start + 2}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 2}:J{r_cust_start + 2}')
    ws[f'I{r_cust_start + 2}'] = invoice_date
    ws[f'I{r_cust_start + 2}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 2}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 3}'] = "Place Of Supply:"
    ws[f'G{r_cust_start + 3}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 3}:J{r_cust_start + 3}')
    ws[f'I{r_cust_start + 3}'] = "Andhra Pradesh"
    ws[f'I{r_cust_start + 3}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 3}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 4}'] = "Corteva PO No :"
    ws[f'G{r_cust_start + 4}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 4}:J{r_cust_start + 4}')
    ws[f'I{r_cust_start + 4}'] = po_number
    ws[f'I{r_cust_start + 4}'].font = styles['font_green_bold']
    ws[f'I{r_cust_start + 4}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 5}'] = "Product:"
    ws[f'G{r_cust_start + 5}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 5}:J{r_cust_start + 5}')
    ws[f'I{r_cust_start + 5}'] = metadata.get('product', '')
    ws[f'I{r_cust_start + 5}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 5}'].alignment = styles['align_center']

    ws[f'G{r_cust_start + 6}'] = "Crop:"
    ws[f'G{r_cust_start + 6}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 6}:J{r_cust_start + 6}')
    ws[f'I{r_cust_start + 6}'] = metadata.get('crop', '')
    ws[f'I{r_cust_start + 6}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 6}'].alignment = styles['align_center']

    # Row AREA
    ws[f'G{r_cust_start + 7}'] = "AREA"
    ws[f'G{r_cust_start + 7}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 7}:J{r_cust_start + 7}')
    ws[f'I{r_cust_start + 7}'] = metadata.get('area') or metadata.get('territory', '')
    ws[f'I{r_cust_start + 7}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 7}'].alignment = styles['align_center']

    # Row ZDGM
    ws[f'G{r_cust_start + 8}'] = "ZDGM"
    ws[f'G{r_cust_start + 8}'].font = styles['font_bold']
    ws.merge_cells(f'I{r_cust_start + 8}:J{r_cust_start + 8}')
    ws[f'I{r_cust_start + 8}'] = metadata.get('zdgm', '') or metadata.get('amm', '')
    ws[f'I{r_cust_start + 8}'].font = styles['font_bold']
    ws[f'I{r_cust_start + 8}'].alignment = styles['align_center']

    # Spacer (Row 23 / header gap)
    r_hdr_spacer = r_cust_start + 9
    ws.row_dimensions[r_hdr_spacer].height = 27

    # Particulars Table Header
    r_tbl_hdr = r_hdr_spacer + 1
    ws.row_dimensions[r_tbl_hdr].height = 20
    ws[f'A{r_tbl_hdr}'] = "SI.NO"
    ws[f'A{r_tbl_hdr}'].font = styles['font_header']
    ws[f'A{r_tbl_hdr}'].alignment = styles['align_center']

    ws.merge_cells(f'B{r_tbl_hdr}:G{r_tbl_hdr}')
    ws[f'B{r_tbl_hdr}'] = "PARTICULARS"
    ws[f'B{r_tbl_hdr}'].font = styles['font_header']
    ws[f'B{r_tbl_hdr}'].alignment = styles['align_center']

    ws[f'H{r_tbl_hdr}'] = "HSN/SAC"
    ws[f'H{r_tbl_hdr}'].font = styles['font_header']
    ws[f'H{r_tbl_hdr}'].alignment = styles['align_center']

    ws[f'I{r_tbl_hdr}'] = "NO.OF.Days /Act"
    ws[f'I{r_tbl_hdr}'].font = styles['font_header']
    ws[f'I{r_tbl_hdr}'].alignment = styles['align_center']

    ws[f'J{r_tbl_hdr}'] = "Amount"
    ws[f'J{r_tbl_hdr}'].font = styles['font_header']
    ws[f'J{r_tbl_hdr}'].alignment = styles['align_center']

    for c in range(1, 11):
        ws.cell(r_tbl_hdr, c).border = styles['box_border']

    # Particulars Data Rows (9 rows)
    start_act_r = r_tbl_hdr + 1
    num_act_rows = max(9, len(activity_groups))
    end_act_r = start_act_r + num_act_rows - 1

    curr_r = start_act_r
    subtotal_sum = 0.0
    total_qty = 0

    for idx, (act_name, act_info) in enumerate(activity_groups.items(), start=1):
        ws.row_dimensions[curr_r].height = 15
        ws.cell(curr_r, 1, idx).alignment = styles['align_center']
        ws.cell(curr_r, 1).font = styles['font_bold']
        
        ws.merge_cells(start_row=curr_r, start_column=2, end_row=curr_r, end_column=7)
        ws.cell(curr_r, 2, f"{act_name} Activities Expenses").font = styles['font_bold']
        ws.cell(curr_r, 2).alignment = styles['align_left']

        if idx == 1:
            ws.cell(curr_r, 8, "998596").alignment = styles['align_center']
            ws.cell(curr_r, 8).font = styles['font_bold']

        q = act_info.get('qty', len(act_info.get('rows', [])))
        total_qty += q
        c_q = ws.cell(curr_r, 9, q)
        c_q.alignment = styles['align_center']
        c_q.font = styles['font_bold']

        raw_amt = act_info.get('raw_amount', 0.0)
        amt_with_sc = raw_amt * (1 + service_charge_pct / 100.0)
        subtotal_sum += amt_with_sc

        c_amt = ws.cell(curr_r, 10, amt_with_sc)
        c_amt.alignment = styles['align_right']
        c_amt.font = styles['font_bold']
        c_amt.number_format = '#,##0.00'

        for c in range(1, 11):
            ws.cell(curr_r, c).border = styles['box_border']

        curr_r += 1

    # Fill remaining empty rows up to end_act_r
    while curr_r <= end_act_r:
        ws.row_dimensions[curr_r].height = 15
        ws.merge_cells(start_row=curr_r, start_column=2, end_row=curr_r, end_column=7)
        
        c_q = ws.cell(curr_r, 9, 0)
        c_q.alignment = styles['align_center']
        c_q.font = styles['font_bold']

        c_amt = ws.cell(curr_r, 10, 0.0)
        c_amt.alignment = styles['align_right']
        c_amt.font = styles['font_bold']
        c_amt.number_format = '#,##0.00'

        for c in range(1, 11):
            ws.cell(curr_r, c).border = styles['box_border']
        curr_r += 1

    # Total Days / Subtotal Row
    tot_days_r = end_act_r + 1
    ws.row_dimensions[tot_days_r].height = 18
    ws.merge_cells(start_row=tot_days_r, start_column=2, end_row=tot_days_r, end_column=7)

    c_tot_q = ws.cell(tot_days_r, 9, f"=SUM(I{start_act_r}:I{end_act_r})")
    c_tot_q.alignment = styles['align_center']
    c_tot_q.font = styles['font_bold']

    c_tot_amt = ws.cell(tot_days_r, 10, f"=SUM(J{start_act_r}:J{end_act_r})")
    c_tot_amt.alignment = styles['align_right']
    c_tot_amt.font = styles['font_bold']
    c_tot_amt.number_format = '#,##0.00'

    for c in range(1, 11):
        ws.cell(tot_days_r, c).border = Border(
            top=Side(style='medium', color='000000'),
            bottom=Side(style='medium', color='000000'),
            left=Side(style='thin', color='000000') if c in [1, 8, 9, 10] else None,
            right=Side(style='thin', color='000000') if c in [7, 8, 9, 10] else None
        )

    # Taxes & Grand Total Block
    r_subtotal = tot_days_r + 1
    ws.row_dimensions[r_subtotal].height = 16
    ws.merge_cells(f'G{r_subtotal}:H{r_subtotal}')
    ws[f'G{r_subtotal}'] = "Sub Total"
    ws[f'G{r_subtotal}'].font = styles['font_bold']
    ws[f'G{r_subtotal}'].alignment = styles['align_left']
    ws[f'J{r_subtotal}'] = f"=J{tot_days_r}"
    ws[f'J{r_subtotal}'].font = styles['font_bold']
    ws[f'J{r_subtotal}'].alignment = styles['align_right']
    ws[f'J{r_subtotal}'].number_format = '#,##0.00'

    r_cgst = r_subtotal + 1
    ws.row_dimensions[r_cgst].height = 16
    ws[f'G{r_cgst}'] = "CGST"
    ws[f'G{r_cgst}'].font = styles['font_bold']
    ws.merge_cells(f'H{r_cgst}:I{r_cgst}')
    ws[f'H{r_cgst}'] = "9%"
    ws[f'H{r_cgst}'].font = styles['font_bold']
    ws[f'H{r_cgst}'].alignment = styles['align_center']
    ws[f'J{r_cgst}'] = f"=ROUND(J{r_subtotal}*0.09, 2)"
    ws[f'J{r_cgst}'].font = styles['font_bold']
    ws[f'J{r_cgst}'].alignment = styles['align_right']
    ws[f'J{r_cgst}'].number_format = '#,##0.00'

    r_sgst = r_subtotal + 2
    ws.row_dimensions[r_sgst].height = 16
    ws[f'G{r_sgst}'] = "SGST"
    ws[f'G{r_sgst}'].font = styles['font_bold']
    ws.merge_cells(f'H{r_sgst}:I{r_sgst}')
    ws[f'H{r_sgst}'] = "9%"
    ws[f'H{r_sgst}'].font = styles['font_bold']
    ws[f'H{r_sgst}'].alignment = styles['align_center']
    ws[f'J{r_sgst}'] = f"=ROUND(J{r_subtotal}*0.09, 2)"
    ws[f'J{r_sgst}'].font = styles['font_bold']
    ws[f'J{r_sgst}'].alignment = styles['align_right']
    ws[f'J{r_sgst}'].number_format = '#,##0.00'

    r_grand = r_subtotal + 3
    ws.row_dimensions[r_grand].height = 18
    ws[f'G{r_grand}'] = "Grand Total"
    ws[f'G{r_grand}'].font = styles['font_bold']
    ws.merge_cells(f'H{r_grand}:I{r_grand}')
    ws[f'H{r_grand}'] = "(Rounded Off)"
    ws[f'H{r_grand}'].font = styles['font_bold']
    ws[f'H{r_grand}'].alignment = styles['align_center']
    ws[f'J{r_grand}'] = f"=ROUND(SUM(J{r_subtotal}:J{r_sgst}), 0)"
    ws[f'J{r_grand}'].font = styles['font_bold']
    ws[f'J{r_grand}'].alignment = styles['align_right']
    ws[f'J{r_grand}'].number_format = '#,##0.00'

    for r in range(r_subtotal, r_grand + 1):
        for c in range(1, 11):
            ws.cell(r, c).border = styles['box_border']

    # Amount in words row
    r_words = r_grand + 1
    ws.row_dimensions[r_words].height = 18
    ws.merge_cells(f'A{r_words}:J{r_words}')
    cgst_calc = round(subtotal_sum * 0.09, 2)
    sgst_calc = round(subtotal_sum * 0.09, 2)
    grand_tot_calc = round(subtotal_sum + cgst_calc + sgst_calc)
    words = num_to_indian_words(grand_tot_calc)
    ws[f'A{r_words}'] = words
    ws[f'A{r_words}'].font = styles['font_bold']
    ws[f'A{r_words}'].alignment = styles['align_center']
    for c in range(1, 11):
        ws.cell(r_words, c).border = styles['box_border']

    # Bank Details & Signatures attached directly after words row (no gap)
    start_bank_r = r_words + 1
    r_bank_end = write_bank_and_signature_block(ws, start_bank_r, styles)

    return r_bank_end, r_subtotal, r_grand


def build_corteva_sheet1_invoice(ws, invoice_no, invoice_date, po_number, metadata, activity_groups, service_charge_pct, styles):
    """
    Builds Sheet1 Tax Invoice for Corteva with 5 rows gap at top, both ORIGINAL and DUPLICATE copies, and balanced margins.
    """
    ws.views.sheetView[0].showGridLines = True

    # Exact column widths calibrated for 1-page A4 Portrait with balanced side margins
    col_widths = {'A': 5.5, 'B': 10.5, 'C': 10.5, 'D': 9.5, 'E': 9.5, 'F': 9.5, 'G': 12.5, 'H': 10.5, 'I': 13.5, 'J': 14.5}
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w

    # 5 top gap rows for Page 1
    for r_top in range(1, 6):
        ws.row_dimensions[r_top].height = 14

    # 1. Render ORIGINAL Invoice block starting at Row 6
    orig_start_r = 6
    orig_end_r, r_subtotal, r_grand = render_corteva_invoice_block(
        ws, orig_start_r, "ORIGINAL", invoice_no, invoice_date, po_number, metadata, activity_groups, service_charge_pct, styles
    )

    # 2. Insert Page Break right after ORIGINAL invoice block
    ws.row_breaks.append(Break(id=orig_end_r))

    # 5 top gap rows for Page 2
    for r_top2 in range(orig_end_r + 1, orig_end_r + 6):
        ws.row_dimensions[r_top2].height = 14

    # 3. Render DUPLICATE Invoice block starting after 5 gap rows
    dup_start_r = orig_end_r + 6
    dup_end_r, _, _ = render_corteva_invoice_block(
        ws, dup_start_r, "DUPLICATE", invoice_no, invoice_date, po_number, metadata, activity_groups, service_charge_pct, styles
    )

    # Set Print Area spanning both ORIGINAL and DUPLICATE pages with balanced side margins
    setup_page_print_fit(ws, print_area=f"A1:J{dup_end_r}", orientation="portrait")

    return r_subtotal, r_grand


def build_corteva_sheet2_details(ws, short_iv, records, service_charge_pct, styles):
    """
    Builds Sheet2 for Corteva matching SS2 layout with PO Number column included:
    Grouped by Activity, then by TBM, with IV NO : {IV} ( index ), subtotals, and 5% service charge row.
    """
    ws.views.sheetView[0].showGridLines = True
    setup_page_print_fit(ws, print_area=None, orientation="landscape")

    col_widths = {
        'A': 6, 'B': 12, 'C': 16, 'D': 16, 'E': 14, 'F': 14, 'G': 14, 'H': 12,
        'I': 12, 'J': 14, 'K': 12, 'L': 15, 'M': 14, 'N': 12, 'O': 12, 'P': 14, 'Q': 20
    }
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w

    act_groups = {}
    for r in records:
        act = clean_str(r.get('activity', '')).upper() or "GENERAL"
        tbm = clean_str(r.get('tbm', '')).title() or "TBM"
        terr = clean_str(r.get('territory', '')).title() or ""
        
        if act not in act_groups:
            act_groups[act] = {}
        tbm_key = (tbm, terr)
        if tbm_key not in act_groups[act]:
            act_groups[act][tbm_key] = []
        act_groups[act][tbm_key].append(r)

    current_r = 1
    iv_idx = 1
    summary_items = []

    headers = [
        "S.No", "Date", "ZDGM", "TBM", "MDO", "Territory", "Product", "Crop", "Activity", "Village",
        "No.of Farmers", "Tent/Hall Suppliers Charges", "Food Expenses", "Transport", "Others", "Total", "PO Number"
    ]

    for act_name, tbm_dict in act_groups.items():
        # Identify Product and Crop for this activity group
        prod_val = ""
        crop_val = ""
        for rows_list in tbm_dict.values():
            for row_item in rows_list:
                if not prod_val and row_item.get('product'):
                    prod_val = str(row_item.get('product')).strip()
                if not crop_val and row_item.get('crop'):
                    crop_val = str(row_item.get('crop')).strip()

        # Section Header: IV NO : 67 ( 1 ) across Columns 1 to 17 (A to Q)
        ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=17)
        c_iv = ws.cell(current_r, 1, f"IV NO : {short_iv} ( {iv_idx} )")
        c_iv.font = styles['font_green_title']
        c_iv.alignment = styles['align_center']
        ws.row_dimensions[current_r].height = 22
        current_r += 1
        iv_idx += 1

        act_table_tot_rows = []

        for (tbm_name, terr_name), rows_list in tbm_dict.items():
            ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=17)
            terr_str = f"{terr_name} " if terr_name else ""
            c_tbl = ws.cell(current_r, 1, f"Activities Expenses by {terr_str}Tbm {tbm_name}")
            c_tbl.font = styles['font_green_title']
            c_tbl.alignment = styles['align_center']
            ws.row_dimensions[current_r].height = 20
            current_r += 1

            hdr_r = current_r
            ws.row_dimensions[hdr_r].height = 24
            for c_idx, h in enumerate(headers, start=1):
                cell = ws.cell(hdr_r, c_idx, h)
                cell.font = styles['font_green_bold'] if c_idx in [1, 9, 10, 11, 12, 13, 14, 15, 16, 17] else styles['font_bold']
                cell.alignment = styles['align_center_wrap']
                cell.border = styles['thin_border']
            current_r += 1

            start_data_r = current_r
            for s_no, item in enumerate(rows_list, start=1):
                ws.row_dimensions[current_r].height = 18
                ws.cell(current_r, 1, s_no).alignment = styles['align_center']
                ws.cell(current_r, 1).font = styles['font_bold']

                ws.cell(current_r, 2, item.get('date', '')).alignment = styles['align_center']
                ws.cell(current_r, 3, item.get('zdgm', '')).alignment = styles['align_left']
                ws.cell(current_r, 4, item.get('tbm', '')).alignment = styles['align_left']
                ws.cell(current_r, 5, item.get('mdo', '')).alignment = styles['align_left']
                ws.cell(current_r, 6, item.get('territory', '')).alignment = styles['align_center']
                ws.cell(current_r, 7, item.get('product', '')).alignment = styles['align_center']
                ws.cell(current_r, 8, item.get('crop', '')).alignment = styles['align_center']
                ws.cell(current_r, 9, item.get('activity', '')).alignment = styles['align_center']
                ws.cell(current_r, 10, item.get('village', '')).alignment = styles['align_left']
                
                c_farm = ws.cell(current_r, 11, item.get('farmers', 0))
                c_farm.alignment = styles['align_center']
                c_farm.number_format = '#,##0'

                for c_off, key in enumerate(['tent', 'food', 'transport', 'others'], start=12):
                    v = item.get(key, 0.0)
                    cell_v = ws.cell(current_r, c_off, v if v > 0 else "")
                    cell_v.alignment = styles['align_right']
                    if v > 0: cell_v.number_format = '#,##0'

                c_row_tot = ws.cell(current_r, 16, f"=SUM(L{current_r}:O{current_r})")
                c_row_tot.alignment = styles['align_right']
                c_row_tot.font = styles['font_bold']
                c_row_tot.number_format = '#,##0'

                # Col 17 (Q): Full PO Number
                c_po = ws.cell(current_r, 17, item.get('po_number', ''))
                c_po.alignment = styles['align_center']
                c_po.font = styles['font_bold']

                for c in range(1, 18):
                    ws.cell(current_r, c).border = styles['thin_border']
                current_r += 1

            end_data_r = current_r - 1

            tot_r = current_r
            ws.row_dimensions[tot_r].height = 20
            ws.cell(tot_r, 15, "Total").font = styles['font_green_bold']
            ws.cell(tot_r, 15).alignment = styles['align_right']
            
            c_tbl_sum = ws.cell(tot_r, 16, f"=SUM(P{start_data_r}:P{end_data_r})")
            c_tbl_sum.font = styles['font_green_bold']
            c_tbl_sum.alignment = styles['align_right']
            c_tbl_sum.number_format = '#,##0'

            for c in range(1, 18):
                ws.cell(tot_r, c).border = styles['thin_border']
            
            act_table_tot_rows.append(tot_r)
            current_r += 1

        subtot_r = current_r
        ws.row_dimensions[subtot_r].height = 20
        ws.merge_cells(start_row=subtot_r, start_column=8, end_row=subtot_r, end_column=9)
        ws.cell(subtot_r, 8, f"Total {act_name}").font = styles['font_bold']
        ws.cell(subtot_r, 8).alignment = styles['align_center']
        
        sum_refs = [f"P{r}" for r in act_table_tot_rows]
        c_sub_val = ws.cell(subtot_r, 10, f"=SUM({','.join(sum_refs)})")
        c_sub_val.font = styles['font_bold']
        c_sub_val.alignment = styles['align_right']
        c_sub_val.number_format = '#,##0.00'

        sc_r = current_r + 1
        ws.row_dimensions[sc_r].height = 18
        c_sc_val = ws.cell(sc_r, 10, f"=J{subtot_r}*{service_charge_pct/100.0:.4f}")
        c_sc_val.font = styles['font_bold']
        c_sc_val.alignment = styles['align_right']
        c_sc_val.number_format = '#,##0.00'

        tot_sc_r = current_r + 2
        ws.row_dimensions[tot_sc_r].height = 20
        c_tot_sc = ws.cell(tot_sc_r, 10, f"=J{subtot_r}+J{sc_r}")
        c_tot_sc.font = styles['font_bold']
        c_tot_sc.alignment = styles['align_right']
        c_tot_sc.number_format = '#,##0.00'

        for r_box in range(subtot_r, tot_sc_r + 1):
            for c in range(1, 18):
                ws.cell(r_box, c).border = styles['thin_border']

        summary_items.append({
            'product': prod_val,
            'crop': crop_val,
            'activity': act_name,
            'subtot_ref': f"J{subtot_r}",
            'sc_ref': f"J{sc_r}",
            'tot_sc_ref': f"J{tot_sc_r}"
        })

        current_r = tot_sc_r + 2

    # Render Cumulative Summary Table at the end of Sheet2
    if summary_items:
        current_r = render_details_cumulative_summary_table(
            ws, start_row=current_r + 1, summary_items=summary_items, service_charge_pct=service_charge_pct, styles=styles
        )


def build_corteva_summary_sheet(ws, invoice_no, invoice_date, po_number, metadata, po_value, r_subtotal, r_grand, styles):
    """
    Builds the 11-column Corteva Invoice Summary sheet dynamically linking to Sheet1 Sub Total & Grand Total.
    """
    ws.views.sheetView[0].showGridLines = True
    setup_page_print_fit(ws, print_area=None, orientation="landscape")

    col_widths = {
        'A': 16, 'B': 16, 'C': 28, 'D': 24, 'E': 18, 'F': 16,
        'G': 24, 'H': 24, 'I': 18, 'J': 18, 'K': 32
    }
    for col_let, w in col_widths.items():
        ws.column_dimensions[col_let].width = w

    headers = [
        "Vendor Code", "Payment Term", "Entity", "Vendor Name", "Invoice No.",
        "Invoice Date", "Invoice Amount (EXC GST)", "Invoice Amount (INC GST)",
        "PO Number", "PO Value", "Name of Corteva Requester/Receiver"
    ]
    
    ws.row_dimensions[1].height = 26
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(1, c_idx, h)
        cell.font = styles['font_header']
        cell.alignment = styles['align_center_wrap']
        cell.border = styles['box_border']

    # Row 2 Data (Linking to Sheet1 J{r_subtotal} and J{r_grand})
    ws.row_dimensions[2].height = 22
    ws.cell(2, 1, "80141626").alignment = styles['align_center']
    ws.cell(2, 2, "45 Days").alignment = styles['align_center']
    ws.cell(2, 3, "Corteva Agriscience India Private Limited").alignment = styles['align_left']
    ws.cell(2, 4, "Radhadevi Kamisetty").alignment = styles['align_left']
    ws.cell(2, 5, invoice_no).alignment = styles['align_center']
    ws.cell(2, 6, invoice_date).alignment = styles['align_center']
    
    # Amount Exc GST linking to Sheet1 Sub Total
    c_exc = ws.cell(2, 7, f"=Sheet1!J{r_subtotal}")
    c_exc.alignment = styles['align_right']
    c_exc.number_format = '#,##0'

    # Amount Inc GST linking to Sheet1 Grand Total
    c_inc = ws.cell(2, 8, f"=Sheet1!J{r_grand}")
    c_inc.alignment = styles['align_right']
    c_inc.number_format = '#,##0'

    ws.cell(2, 9, po_number).alignment = styles['align_center']
    
    c_val = ws.cell(2, 10, po_value if po_value else 250000)
    c_val.alignment = styles['align_right']
    c_val.number_format = '#,##0'

    requester = metadata.get('zdgm', '') or metadata.get('amm', '') or "R.Bhaskar"
    ws.cell(2, 11, requester).alignment = styles['align_left']

    for c in range(1, 12):
        ws.cell(2, c).font = styles['font_regular']
        ws.cell(2, c).border = styles['box_border']
